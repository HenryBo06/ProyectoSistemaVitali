"""Native Community accounting check. Run in odoo-bin shell; fixtures roll back."""
import calendar
import csv
import hashlib
import json
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4

from odoo import fields
from odoo.exceptions import UserError


assert env.cr.dbname == "vitali_lab", "Synthetic laboratory only"
root = Path.cwd()
output = root / ".local-odoo/evidence/accounting"
tag = "LAB-CLOSE-" + uuid4().hex[:10]
lab_company = env.company
today = fields.Date.context_today(env["res.company"].with_context(tz="America/El_Salvador"))
period_start = today.replace(day=1)
period_end = today.replace(day=calendar.monthrange(today.year, today.month)[1])
results = []
reports = []
kg = env.ref("uom.product_uom_kgm")
config = env["ir.config_parameter"].sudo()
seed_sale_id = int(config.get_param("vitali_lab.sale_id"))
seed_mo_id = int(config.get_param("vitali_lab.manufacturing_id"))


def close(actual, expected):
    assert abs(actual - expected) < .001, (actual, expected)


def passed(scenario, **values):
    results.append({"scenario": scenario, "status": "passed", **values})
    print("PASS:", scenario, json.dumps(values, ensure_ascii=True))


def preserved():
    sale = env["sale.order"].browse(seed_sale_id)
    mo = env["mrp.production"].browse(seed_mo_id)
    integrated = env["sale.order"].browse(11).exists()
    return {"sale": [sale.id, sale.state, sale.order_line.product_uom_qty, sale.invoice_ids.ids],
            "manufacturing": [mo.id, mo.state, mo.product_qty],
            "integrated": [integrated.name, integrated.order_line.product_uom_qty,
                integrated.order_line.qty_delivered,
                [(m.id, m.state, m.amount_total, m.amount_residual) for m in integrated.invoice_ids]]
                if integrated else None}


def statement(company, scope):
    """One owning aggregation of posted journal items, with period and closing balances."""
    lines = company.env["account.move.line"].search([
        ("company_id", "=", company.id), ("parent_state", "=", "posted"),
        ("date", "<=", period_end)], order="date,move_id,id")
    accounts = defaultdict(lambda: {"opening": 0., "debit": 0., "credit": 0., "balance": 0.})
    ledger = []
    for line in lines:
        row = accounts[line.account_id]
        row["balance"] += line.balance
        if line.date < period_start:
            row["opening"] += line.balance
        else:
            row["debit"] += line.debit
            row["credit"] += line.credit
            ledger.append({"scope": scope, "line_id": line.id, "move_id": line.move_id.id,
                "document": line.move_id.name, "date": str(line.date),
                "journal": line.journal_id.name, "account_code": line.account_id.code,
                "account": line.account_id.name, "account_type": line.account_id.account_type,
                "description": line.name or "", "debit": round(line.debit, 2),
                "credit": round(line.credit, 2), "balance": round(line.balance, 2)})
    trial = [{"scope": scope, "account_code": account.code, "account": account.name,
        "account_type": account.account_type, **{k: round(v, 2) for k, v in totals.items()}}
        for account, totals in sorted(accounts.items(), key=lambda item: item[0].code)]
    assets = sum(v["balance"] for a, v in accounts.items() if a.account_type.startswith("asset"))
    liabilities = -sum(v["balance"] for a, v in accounts.items() if a.account_type.startswith("liability"))
    equity = -sum(v["balance"] for a, v in accounts.items() if a.account_type.startswith("equity"))
    cumulative_profit = -sum(v["balance"] for a, v in accounts.items()
                            if a.account_type.startswith(("income", "expense")))
    income = sum(v["credit"] - v["debit"] for a, v in accounts.items()
                  if a.account_type.startswith("income"))
    expenses = sum(v["debit"] - v["credit"] for a, v in accounts.items()
                   if a.account_type.startswith("expense"))
    close(sum(row["debit"] for row in trial), sum(row["credit"] for row in trial))
    close(assets, liabilities + equity + cumulative_profit)
    summary = {"scope": scope, "company": company.name, "currency": company.currency_id.name,
        "period_start": str(period_start), "period_end": str(period_end),
        "posted_move_count": len({row["move_id"] for row in ledger}), "posted_line_count": len(ledger),
        "assets": round(assets, 2), "liabilities": round(liabilities, 2), "equity": round(equity, 2),
        "cumulative_profit": round(cumulative_profit, 2), "income": round(income, 2),
        "expenses": round(expenses, 2), "period_profit": round(income - expenses, 2),
        "debit": round(sum(row["debit"] for row in trial), 2),
        "credit": round(sum(row["credit"] for row in trial), 2)}
    return {"summary": summary, "trial_balance": trial, "ledger": ledger}


def finish(picking):
    picking.move_ids.write({"picked": True})
    action = picking.button_validate()
    if isinstance(action, dict) and action.get("res_model") == "stock.backorder.confirmation":
        picking.env[action["res_model"]].with_context(**action.get("context", {})).create({}).process()
    assert picking.state == "done", (picking.name, picking.state, action)


seed_before = preserved()
try:
    # ponytail: a temporary company isolates native closing/account properties; rollback is the cleanup.
    company = env["res.company"].create({"name": tag + " simulated accounting",
        "currency_id": env.ref("base.USD").id, "country_id": env.ref("base.sv").id})
    test = env(context=dict(env.context, allowed_company_ids=company.ids,
                            tz="America/El_Salvador", lang="es_ES"))
    company = company.with_env(test)
    test["account.chart.template"].try_loading("generic_coa", company, install_demo=False)
    company.write({"inventory_valuation": "periodic", "inventory_period": "manual"})
    warehouse = test["stock.warehouse"].search([("company_id", "=", company.id)], limit=1)
    if not warehouse:
        warehouse = test["stock.warehouse"].create({"name": tag + " warehouse", "code": "CLOSE",
                                                    "company_id": company.id})
    category = test["product.category"].create({"name": tag,
        "property_valuation": "periodic", "property_cost_method": "standard",
        "property_stock_valuation_account_id": company.account_stock_valuation_id.id,
        "property_stock_journal": company.account_stock_journal_id.id,
        "property_account_expense_categ_id": company.expense_account_id.id,
        "property_account_income_categ_id": company.income_account_id.id})
    variation = company.account_stock_valuation_id.account_stock_variation_id
    assert company.account_stock_journal_id and company.account_stock_valuation_id and variation
    raw = test["product.product"].create({"name": tag + " raw", "default_code": tag + "-MP",
        "type": "consu", "is_storable": True, "company_id": company.id,
        "uom_id": kg.id, "categ_id": category.id, "standard_price": 2,
        "taxes_id": [(5, 0, 0)], "supplier_taxes_id": [(5, 0, 0)]})
    finished_category = category.copy({"name": tag + " manufactured", "property_cost_method": "average"})
    finished_product = test["product.product"].create({"name": tag + " finished",
        "default_code": tag + "-PT", "type": "consu", "is_storable": True,
        "company_id": company.id, "uom_id": kg.id, "categ_id": finished_category.id,
        "standard_price": 0, "list_price": 4,
        "taxes_id": [(5, 0, 0)], "supplier_taxes_id": [(5, 0, 0)]})
    supplier = test["res.partner"].create({"name": tag + " supplier", "company_id": company.id,
                                          "supplier_rank": 1})
    customer = test["res.partner"].create({"name": tag + " customer", "company_id": company.id,
                                          "customer_rank": 1})
    cash = test["account.journal"].create({"name": tag + " cash", "code": "CCASH",
                                           "type": "cash", "company_id": company.id})
    (cash.inbound_payment_method_line_ids | cash.outbound_payment_method_line_ids).write({
        "payment_account_id": cash.default_account_id.id})
    general = company.account_stock_journal_id
    equity = test["account.account"].search([("company_ids", "in", company.ids),
                                               ("account_type", "=", "equity")], limit=1)
    opening = test["account.move"].create({"journal_id": general.id, "date": today,
        "ref": tag + " capital", "line_ids": [(0, 0, {"account_id": cash.default_account_id.id,
            "debit": 100}), (0, 0, {"account_id": equity.id, "credit": 100})]})
    opening.action_post()
    purchase = test["purchase.order"].create({"partner_id": supplier.id, "company_id": company.id,
        "picking_type_id": warehouse.in_type_id.id, "order_line": [(0, 0, {"product_id": raw.id,
            "product_qty": 10, "product_uom_id": kg.id, "price_unit": 2})]})
    purchase.button_confirm()
    receipt = purchase.picking_ids
    receipt.move_ids.quantity = 10
    finish(receipt)
    purchase.action_create_invoice()
    bill = purchase.invoice_ids
    bill.write({"invoice_date": today, "date": today})
    bill.action_post()
    payment = test["account.payment.register"].with_context(active_model="account.move",
        active_ids=bill.ids).create({"amount": 8, "payment_date": today, "journal_id": cash.id,
        "payment_method_line_id": cash.outbound_payment_method_line_ids[:1].id})._create_payments()
    close(bill.amount_total, 20)
    close(bill.amount_residual, 12)
    payable = bill.line_ids.filtered(lambda line: line.account_id.account_type == "liability_payable")
    assert payable.matched_debit_ids and not payable.reconciled
    passed("supplier_bill_partial_payment_native_reconciliation", bill_usd=20, paid_usd=8,
        payable_usd=12, bill_state=bill.state, payment_state=bill.payment_state,
        partial_reconcile_ids=payable.matched_debit_ids.ids)

    original_partial_ids = payable.matched_debit_ids.ids
    # ponytail: native savepoint isolates this credit check from the existing closing figures.
    with test.cr.savepoint() as credit_savepoint:
        supplier_credit = bill._reverse_moves(default_values_list=[{
            "ref": tag + " supplier return credit", "invoice_date": today, "date": today}], cancel=False)
        credit_line = supplier_credit.invoice_line_ids.filtered(lambda line: line.product_id == raw)
        original_line = bill.invoice_line_ids.filtered(lambda line: line.product_id == raw)
        assert len(credit_line) == len(original_line) == 1
        credit_line.quantity = 2
        supplier_credit.action_post()
        assert supplier_credit.move_type == "in_refund" and supplier_credit.state == "posted"
        assert supplier_credit.reversed_entry_id == bill and supplier_credit in bill.reversal_move_ids
        assert credit_line.account_id == original_line.account_id == company.expense_account_id
        close(supplier_credit.amount_total, 4)
        close(credit_line.credit, 4)
        close(credit_line.debit, 0)
        credit_payable = supplier_credit.line_ids.filtered(lambda line:
            line.account_id.account_type == "liability_payable")
        assert credit_payable.account_id == payable.account_id
        close(credit_payable.debit, 4)
        close(credit_payable.credit, 0)
        (payable | credit_payable).reconcile()
        close(bill.amount_residual, 8)
        close(supplier_credit.amount_residual, 0)
        assert not payable.reconciled and credit_payable.reconciled
        reconciliations = payable.matched_debit_ids.filtered(lambda item:
            item.debit_move_id == credit_payable and item.credit_move_id == payable)
        assert len(reconciliations) == 1
        close(reconciliations.amount, 4)
        credit_evidence = {"bill_usd": 20, "partial_payment_usd": 8, "payable_before_usd": 12,
            "credited_kg": 2, "price_per_kg_usd": 2, "supplier_credit_usd": 4,
            "payable_after_usd": 8, "credit_residual_usd": 0, "move_type": supplier_credit.move_type,
            "credit_move_id": supplier_credit.id, "original_bill_id": bill.id,
            "original_line_id": original_line.id, "credit_line_id": credit_line.id,
            "reversed_entry_id": supplier_credit.reversed_entry_id.id,
            "partial_reconcile_id": reconciliations.id,
            "expense_credit_usd": 4, "payable_debit_usd": 4,
            "physical_supplier_return": "verified_in_separate_native_operational_case_not_this_purchase"}
        credit_savepoint.rollback()
    close(bill.amount_residual, 12)
    assert payable.matched_debit_ids.ids == original_partial_ids
    assert not test["account.move"].browse(credit_evidence["credit_move_id"]).exists()
    passed("supplier_credit_native_reversal_partial_reconciliation_and_savepoint_rollback",
           **credit_evidence, payable_after_rollback_usd=12, credit_removed_after_rollback=True)

    bom = test["mrp.bom"].create({"product_tmpl_id": finished_product.product_tmpl_id.id,
        "company_id": company.id, "product_qty": 1, "product_uom_id": kg.id,
        "bom_line_ids": [(0, 0, {"product_id": raw.id, "product_qty": 1.15,
                                 "product_uom_id": kg.id})]})
    mo = test["mrp.production"].create({"product_id": finished_product.id, "product_qty": 4,
        "product_uom_id": kg.id, "bom_id": bom.id, "company_id": company.id,
        "picking_type_id": warehouse.manu_type_id.id, "extra_cost": .2, "origin": tag})
    mo.action_confirm()
    mo.action_assign()
    mo.qty_producing = 4
    mo._set_qty_producing()
    mo.move_raw_ids.picked = True
    mo.button_mark_done()
    assert mo.state == "done"
    close(sum(mo.move_raw_ids.mapped("value")), 9.2)
    close(sum(mo.move_finished_ids.mapped("value")), 10)
    close(finished_product.standard_price, 2.5)
    overhead = test["account.move"].create({"journal_id": general.id, "date": today,
        "ref": tag + " simulated manufacturing overhead", "line_ids": [
            (0, 0, {"account_id": company.expense_account_id.id, "debit": .8}),
            (0, 0, {"account_id": cash.default_account_id.id, "credit": .8})]})
    overhead.action_post()
    scrap = test["stock.scrap"].create({"product_id": raw.id, "scrap_qty": .2,
        "product_uom_id": kg.id, "company_id": company.id,
        "location_id": warehouse.lot_stock_id.id, "origin": tag})
    scrap.action_validate()
    close(scrap.move_ids.value, .4)
    passed("manufacturing_material_overhead_and_scrap_costs", produced_kg=4,
        raw_consumed_kg=4.6, material_cost_usd=9.2, simulated_overhead_usd=.8,
        manufactured_value_usd=10, unit_cost_usd=2.5, scrap_kg=.2, scrap_value_usd=.4)

    sale = test["sale.order"].create({"partner_id": customer.id, "company_id": company.id,
        "warehouse_id": warehouse.id, "client_order_ref": tag, "order_line": [(0, 0, {
            "product_id": finished_product.id, "product_uom_qty": 4,
            "product_uom_id": kg.id, "price_unit": 4})]})
    sale.action_confirm()
    delivery = sale.picking_ids
    delivery.action_assign()
    delivery.move_ids.quantity = 4
    finish(delivery)
    returning = test["stock.return.picking"].with_context(active_id=delivery.id,
        active_ids=delivery.ids, active_model="stock.picking").create({"picking_id": delivery.id})
    returning.product_return_moves.quantity = 1
    returned = returning._create_return()
    returned.move_ids.quantity = 1
    finish(returned)
    invoice = sale._create_invoices()
    invoice.action_post()
    refund = invoice._reverse_moves(default_values_list=[{"ref": tag + " return"}], cancel=False)
    refund.invoice_line_ids.filtered(lambda line: line.product_id == finished_product).quantity = 1
    refund.action_post()
    (invoice.line_ids | refund.line_ids).filtered(lambda line:
        line.account_id.account_type == "asset_receivable" and not line.reconciled).reconcile()
    for amount in (6, 6):
        test["account.payment.register"].with_context(active_model="account.move",
            active_ids=invoice.ids).create({"amount": amount, "payment_date": today,
            "journal_id": cash.id, "payment_method_line_id": cash.inbound_payment_method_line_ids[:1].id})._create_payments()
    close(invoice.amount_total, 16)
    close(refund.amount_total, 4)
    close(invoice.amount_residual, 0)
    assert invoice.payment_state == "paid" and refund.amount_residual == 0
    close(sale.order_line.qty_delivered, 3)
    passed("sales_return_credit_and_collections", invoice_usd=16, credit_usd=4,
        collection_usd=12, receivable_usd=0, net_delivered_kg=3)

    close(raw.qty_available, 5.2)
    close(finished_product.qty_available, 1)
    close(raw.total_value + finished_product.total_value, 12.9)
    close(sum(company.stock_accounting_value().values()), 0)
    closing_action = company.action_close_stock_valuation(at_date=period_end)
    closing = test["account.move"].browse(closing_action["res_id"])
    closing.action_post()
    assert closing.state == "posted" and closing.date == period_end
    close(sum(company.stock_accounting_value(at_date=period_end).values()), 12.9)
    with test.cr.savepoint():
        try:
            company.action_close_stock_valuation(at_date=period_end)
        except UserError:
            pass
        else:
            raise AssertionError("A repeated balanced inventory closing must not create another entry")
    passed("native_periodic_inventory_closing", native_method="res.company.action_close_stock_valuation",
        raw_inventory_usd=10.4, finished_inventory_usd=2.5, stock_ledger_usd=12.9,
        closing_move_id=closing.id, closing_date=str(closing.date), repeated_closing_created_entry=False)
    excluded_draft = opening.copy({"date": today, "ref": tag + " unposted probe", "name": "/"})
    excluded_next_month = opening.copy({"date": period_end + timedelta(days=1),
        "ref": tag + " outside-period probe", "name": "/"})
    excluded_next_month.action_post()
    fixture_report = statement(company, "cierre_controlado")
    summary = fixture_report["summary"]
    for key, expected in {"assets": 116.1, "liabilities": 12, "equity": 100,
        "income": 12, "expenses": 7.9, "period_profit": 4.1}.items():
        close(summary[key], expected)
    reported_move_ids = {row["move_id"] for row in fixture_report["ledger"]}
    assert closing.id in reported_move_ids and bill.id in reported_move_ids and payment.move_id.id in reported_move_ids
    assert excluded_draft.id not in reported_move_ids and excluded_next_month.id not in reported_move_ids
    reports.append(fixture_report)
    passed("posted_trial_balance_balance_sheet_and_profit_loss", **summary)

    original_lock = company.fiscalyear_lock_date
    excluded_draft.unlink()
    company.fiscalyear_lock_date = period_end
    assert company._get_user_fiscal_lock_date(general) == period_end
    late = opening.copy({"date": today, "ref": tag + " closed-period probe", "name": "/"})
    late.action_post()
    assert late.date > period_end, "Native posting must move outside the closed period"
    with test.cr.savepoint():
        try:
            closing.line_ids.filtered(lambda line: line.debit).write({"account_id": company.income_account_id.id})
        except UserError:
            pass
        else:
            raise AssertionError("Native period lock must reject editing its journal items")
    company.fiscalyear_lock_date = original_lock
    passed("monthly_native_lock_and_restore", field="fiscalyear_lock_date",
        lock_date=str(period_end), attempted_posting_date=str(today), actual_posting_date=str(late.date),
        closed_period_edit_rejected=True, soft_lock_restored=True)
finally:
    env.cr.rollback()
    env.invalidate_all()

assert preserved() == seed_before, "Original and integrated sales must remain unchanged"
assert not env["res.company"].search_count([("name", "=", tag + " simulated accounting")])
passed("rollback_preserves_original_and_integrated_operations", seed_snapshot=seed_before,
       temporary_company_removed=True)
reports.insert(0, statement(lab_company, "laboratorio_actual"))
passed("persisted_lab_posted_financial_statements", **reports[0]["summary"])

output.mkdir(parents=True, exist_ok=True)
exports = {}
for name, rows in {"balanza": [row for report in reports for row in report["trial_balance"]],
                   "asientos": [row for report in reports for row in report["ledger"]],
                   "estados": [report["summary"] for report in reports]}.items():
    path = output / (name + ".csv")
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows({key: "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
                          else value for key, value in row.items()} for row in rows)
    exports[name] = {"path": str(path.relative_to(root)),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "rows": len(rows)}

# ponytail: installed openpyxl provides a reusable native export, no reporting engine or Enterprise module.
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

book = Workbook()
book.remove(book.active)
labels = {"scope": "Escenario", "company": "Empresa", "currency": "Moneda", "period_start": "Inicio",
    "period_end": "Fin", "posted_move_count": "Asientos", "posted_line_count": "Apuntes",
    "assets": "Activos USD", "liabilities": "Pasivos USD", "equity": "Capital USD",
    "cumulative_profit": "Resultado acumulado USD", "income": "Ingresos USD", "expenses": "Gastos USD",
    "period_profit": "Resultado mensual USD", "debit": "Debe USD", "credit": "Haber USD",
    "account_code": "Codigo", "account": "Cuenta", "account_type": "Tipo de cuenta",
    "opening": "Saldo inicial USD", "balance": "Saldo USD", "line_id": "ID apunte", "move_id": "ID asiento",
    "document": "Documento", "date": "Fecha", "journal": "Diario", "description": "Descripcion"}
for title, filename in (("Estados", "estados"), ("Balanza", "balanza"), ("Asientos", "asientos")):
    sheet = book.create_sheet(title)
    sheet.append(["Vitali: cierre contable simulado (USD)"])
    sheet.append([f"{period_start} a {period_end}; fuentes: account.move.line contabilizados en vitali_lab"])
    sheet.append(["laboratorio_actual: datos guardados; cierre_controlado: comprobacion temporal revertida"])
    path = root / exports[filename]["path"]
    with path.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))
    sheet.append([labels[column] for column in rows[0]])
    for row in rows[1:]:
        sheet.append([float(value) if rows[0][col] in {"opening", "debit", "credit", "balance",
            "assets", "liabilities", "equity", "cumulative_profit", "income", "expenses", "period_profit"}
            else int(value) if rows[0][col] in {"line_id", "move_id", "posted_move_count", "posted_line_count"}
            else date.fromisoformat(value) if rows[0][col] in {"period_start", "period_end", "date"}
            else value for col, value in enumerate(row)])
    sheet.freeze_panes = "C5"
    sheet.sheet_view.showGridLines = False
    for row in (1, 2, 3):
        sheet.cell(row, 1).font = Font(name="Calibri", size=15 if row == 1 else 10,
            bold=row == 1, color="167A43" if row == 1 else "555555")
    for cell in sheet[4]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="167A43")
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[4].height = 30
    table = Table(displayName="Table" + title, ref=f"A4:{get_column_letter(sheet.max_column)}{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    sheet.add_table(table)
    for col in range(1, sheet.max_column + 1):
        width = max(len(str(sheet.cell(row, col).value or "")) for row in range(4, sheet.max_row + 1)) + 2
        sheet.column_dimensions[get_column_letter(col)].width = min(48, max(14, width))
        for row in range(5, sheet.max_row + 1):
            cell = sheet.cell(row, col)
            if isinstance(cell.value, (int, float)) and rows[0][col - 1] not in {"line_id", "move_id",
                "posted_move_count", "posted_line_count"}:
                cell.number_format = '#,##0.00;[Red](#,##0.00);"-"'
            elif isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"
    sheet.print_title_rows = "1:4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
path = output / "cierre_contable_simulado.xlsx"
book.save(path)
checked = load_workbook(path, data_only=True)
assert checked.sheetnames == ["Estados", "Balanza", "Asientos"]
assert checked["Asientos"].max_row - 4 == exports["asientos"]["rows"]
headers = [cell.value for cell in checked["Estados"][4]]
close(checked["Estados"].cell(6, headers.index(labels["period_profit"]) + 1).value, 4.1)
exports["xlsx"] = {"path": str(path.relative_to(root)),
    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "sheets": checked.sheetnames}
evidence = {"status": "passed", "generated_utc": datetime.now(timezone.utc).isoformat(),
    "database": "vitali_lab", "source": "posted account.move.line; native stock valuation, payment and locks",
    "odoo_source_commit": "643610127254ee77bbc756b6746739e9e9bb85ff",
    "run_date_el_salvador": str(today),
    "period_start": str(period_start), "period_end": str(period_end), "currency": "USD",
    "synthetic": True, "fiscal_validation": False, "fixtures_rolled_back": True,
    "limits": ["Synthetic generic accounts, no Salvadoran DTE or fiscal validation",
        "Controlled month-end is a simulation, not an actual elapsed month",
        "Persisted lab ledger has no inventory closing or opening capital; its profit is not industrial margin",
        "Manufacturing overhead is a declared USD0.20/kg simulation; labor, energy and transport require real inputs"],
    "checks": results, "reports": reports, "exports": exports}
(root / ".local-odoo/evidence/accounting-close.json").write_text(
    json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print("ACCOUNTING_CLOSE_PASS", json.dumps({"checks": len(results), "exports": exports}, ensure_ascii=True))
