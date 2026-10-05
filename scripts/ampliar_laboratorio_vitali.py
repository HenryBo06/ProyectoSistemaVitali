"""Expand only the synthetic lab: 10 customers, 8 products and coherent history.

Run in the Odoo shell first, then with the SmartOrder Python environment.
Existing seed operations, credentials, source files and sales are preserved.
"""
from datetime import date, timedelta
from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1] if '__file__' in globals() and Path(__file__).name == 'ampliar_laboratorio_vitali.py' else Path.cwd()
CLIENTS = tuple(f'Cliente Vitali LAB-{n:03d} (simulado)' for n in range(1, 11))
PRODUCTS = tuple({'name': 'Producto avícola LAB (simulado)' if n == 1 else f'Producto avícola LAB-{n:03d} (simulado)',
    'sku': f'VIT-LAB-PT-{n:03d}', 'unit': 'kg' if n <= 6 else 'unid',
    'uom': 'uom.product_uom_kgm' if n <= 6 else 'uom.product_uom_unit',
    'price': 4 + (n - 1) / 2, 'cost': 2 if n == 1 else 2 + (n - 1) / 5,
    'raw_per_unit_kg': 1.15, 'pack_per_unit': 0 if n == 1 else 1} for n in range(1, 9))
PORTFOLIOS = {f'lab_vendedor_{n:02d}': CLIENTS[n - 1::3] for n in range(1, 4)}


def prepare_odoo(env):
    from odoo import Command
    assert env.cr.dbname == 'vitali_lab', 'Only the synthetic laboratory may be expanded'
    params = env['ir.config_parameter'].sudo()
    assert params.get_param('vitali_lab.seed_version'), 'Prepare the original laboratory first'
    company = env.company
    warehouse = env['vitali.integration']._operation()._warehouse()
    seed_sale = env['sale.order'].browse(int(params.get_param('vitali_lab.sale_id')))
    seed_mo = env['mrp.production'].browse(int(params.get_param('vitali_lab.manufacturing_id')))
    before = (seed_sale.state, seed_sale.order_line.product_uom_qty, seed_mo.state, seed_mo.product_qty)
    kg, units = env.ref('uom.product_uom_kgm'), env.ref('uom.product_uom_unit')
    admin = env.ref('base.user_admin')
    admin.group_ids = [Command.link(env.ref(xmlid).id) for xmlid in
                       ('mrp.group_mrp_routings', 'stock.group_stock_multi_locations', 'uom.group_uom')]
    pricelist = env['product.pricelist'].search([('currency_id', '=', company.currency_id.id),
        ('company_id', 'in', [False, company.id])], limit=1)
    assert pricelist and company.currency_id.name == 'USD'
    for n, name in enumerate(CLIENTS, 1):
        code = f'VIT-LAB-CLIENT-{n:03d}'
        customer = env['res.partner'].search([('ref', '=', code)])
        assert len(customer) <= 1
        if not customer:
            env['res.partner'].create({'name': name, 'ref': code, 'customer_rank': 1,
                'company_id': company.id, 'lang': 'es_ES', 'property_product_pricelist': pricelist.id})

    # ponytail: native records and explicit test parameters; no extra business model.
    raw = env['product.product'].search([('default_code', '=', 'VIT-LAB-MP-AMPL-001')])
    if not raw:
        raw = env['product.product'].create({'name': 'Insumo ampliado LAB (simulado)',
            'default_code': 'VIT-LAB-MP-AMPL-001', 'type': 'consu', 'is_storable': True,
            'tracking': 'lot', 'uom_id': kg.id, 'standard_price': 1.5,
            'taxes_id': [Command.clear()], 'supplier_taxes_id': [Command.clear()]})
    pack = env['product.product'].search([('default_code', '=', 'VIT-LAB-EMP-001')])
    if not pack:
        pack = env['product.product'].create({'name': 'Empaque LAB (simulado)', 'default_code': 'VIT-LAB-EMP-001',
            'type': 'consu', 'is_storable': True, 'uom_id': units.id, 'standard_price': 0.2,
            'taxes_id': [Command.clear()], 'supplier_taxes_id': [Command.clear()]})
    assert len(raw) == len(pack) == 1
    calendar = env['resource.calendar'].search([('name', '=', 'Calendario LAB continuo (simulado)'), ('company_id', '=', company.id)])
    if not calendar:
        calendar = env['resource.calendar'].create({'name': 'Calendario LAB continuo (simulado)',
            'company_id': company.id, 'tz': 'America/El_Salvador',
            'attendance_ids': [Command.create({'name': f'LAB {day} {period}', 'dayofweek': str(day),
                'day_period': period, 'hour_from': start, 'hour_to': end})
                for day in range(7) for period, start, end in [('morning', 8, 12), ('afternoon', 13, 17)]]})
    center = env['mrp.workcenter'].search([('code', '=', 'VIT-LAB-CENTRO-001')])
    if not center:
        center = env['mrp.workcenter'].create({'name': 'Centro de proceso LAB (simulado)',
            'code': 'VIT-LAB-CENTRO-001', 'company_id': company.id, 'costs_hour': 12,
            'resource_calendar_id': calendar.id,
            'time_start': 5, 'time_stop': 5,
            'note': '<p>SIMULADO: 10 kg o 10 unidades por ciclo de 15 minutos; revisar con responsables.</p>',
            'capacity_ids': [Command.create({'product_uom_id': unit.id, 'capacity': 10}) for unit in (kg, units)]})
    assert len(center) == 1
    center.resource_calendar_id = calendar
    created = []
    for row in PRODUCTS:
        product = env['product.product'].search([('default_code', '=', row['sku'])])
        assert len(product) <= 1
        if not product:
            product = env['product.product'].create({'name': row['name'], 'default_code': row['sku'],
                'type': 'consu', 'is_storable': True, 'tracking': 'lot', 'uom_id': env.ref(row['uom']).id,
                'list_price': row['price'], 'standard_price': row['cost'],
                'taxes_id': [Command.clear()], 'supplier_taxes_id': [Command.clear()]})
            created.append(product)
        assert product.uom_id == env.ref(row['uom'])
        if row['sku'] != 'VIT-LAB-PT-001' and not env['mrp.bom'].search_count([('product_tmpl_id', '=', product.product_tmpl_id.id)]):
            env['mrp.bom'].create({'product_tmpl_id': product.product_tmpl_id.id, 'product_id': product.id,
                'company_id': company.id, 'product_qty': 1, 'product_uom_id': product.uom_id.id,
                'bom_line_ids': [Command.create({'product_id': raw.id, 'product_qty': 1.15, 'product_uom_id': kg.id}),
                                 Command.create({'product_id': pack.id, 'product_qty': 1, 'product_uom_id': units.id})],
                'operation_ids': [Command.create({'name': 'Proceso y empaque LAB (simulado)',
                    'workcenter_id': center.id, 'time_mode': 'manual', 'time_cycle_manual': 15})]})
    locations = []
    for n in range(1, 4):
        name = f'Punto de venta LAB-{n:03d} (simulado)'
        location = env['stock.location'].search([('name', '=', name), ('company_id', '=', company.id)])
        assert len(location) <= 1
        if not location:
            location = env['stock.location'].create({'name': name, 'location_id': warehouse.view_location_id.id,
                'usage': 'internal', 'company_id': company.id})
        locations.append(location.id)
    transit = env['stock.location'].search([('name', '=', 'Tránsito Vitali LAB (simulado)'), ('company_id', '=', company.id)])
    if not transit:
        transit = env['stock.location'].create({'name': 'Tránsito Vitali LAB (simulado)',
            'location_id': warehouse.view_location_id.id, 'usage': 'transit', 'company_id': company.id})
    if not params.get_param('vitali_lab.expanded_catalog_version'):
        for product, quantity in [(p, 20) for p in created] + [(raw, 50), (pack, 200)]:
            lot = env['stock.lot'].create({'name': f'LAB-AMPL-{product.default_code}', 'product_id': product.id,
                                         'company_id': company.id}) if product.tracking == 'lot' else env['stock.lot']
            quant = env['stock.quant'].with_context(inventory_mode=True).create({'product_id': product.id,
                'location_id': warehouse.lot_stock_id.id, 'lot_id': lot.id if lot else False, 'inventory_quantity': quantity})
            quant.action_apply_inventory()
        params.set_param('vitali_lab.expanded_catalog_version', '1')
    after = (seed_sale.state, seed_sale.order_line.product_uom_qty, seed_mo.state, seed_mo.product_qty)
    assert before == after and seed_mo.state == 'draft'
    assert env['res.partner'].search_count([('ref', 'in', [f'VIT-LAB-CLIENT-{n:03d}' for n in range(1, 11)])]) == 10
    assert env['product.product'].search_count([('default_code', 'in', [p['sku'] for p in PRODUCTS])]) == 8
    env.cr.commit()
    evidence = {'status': 'passed', 'customers': 10, 'finished_products': 8, 'central_warehouse': warehouse.name,
        'sales_point_location_ids': locations, 'transit_location_id': transit.id, 'workcenter_id': center.id,
        'new_boms_with_packaging': 7, 'seed_preserved': True, 'products': PRODUCTS,
        'notice': 'All parameters are synthetic. No automatic manufacturing route was enabled.'}
    (ROOT / '.local-odoo/evidence/expanded-catalog.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    print('PASS: native persistent catalog: 10 customers, 8 products, 3 sales points, transit, packaging and capacity')


def prepare_smartorder():
    from preparar_smartorder_lab import LAB_HOME, bootstrap
    bootstrap()
    from django.contrib.auth import get_user_model
    from django.db import transaction
    from openpyxl import Workbook
    from operations.models import Assignment, OdooCustomerMapping, OdooProductMapping, SalesDataset
    from operations.services import import_sales, import_inventory

    marker = LAB_HOME / 'laboratorio.json'
    info = json.loads(marker.read_text(encoding='utf-8'))
    assert info['kind'] == 'SmartOrder synthetic laboratory'
    if info.get('version') == 2:
        assert SalesDataset.objects.get(pk=info['sales_dataset_id']).row_count == info['row_count']
        print('PRESERVED: expanded synthetic history, accounts and native references')
        return
    catalog = json.loads((ROOT / '.local-odoo/evidence/expanded-catalog.json').read_text(encoding='utf-8'))
    assert catalog['status'] == 'passed' and catalog['finished_products'] == 8
    first, last = date.fromisoformat(info['first_date']), date.fromisoformat(info['last_date'])
    folder = LAB_HOME / 'fuentes_simuladas'
    sales_path = folder / 'ventas_SIMULADAS_18_meses_ampliadas.xlsx'
    inventory_path = folder / 'inventario_SIMULADO_ampliado.xlsx'
    ledger_path = folder / 'operacion_SIMULADA_18_meses.xlsx'
    sales, ledger = Workbook(write_only=True), Workbook(write_only=True)
    sheet = sales.create_sheet('Ventas')
    sheet.append(('Fecha', 'Cliente', 'Zona_Geografica', 'Canal_Distribucion', 'Canal_Venta', 'Producto',
                  'Categoria', 'Cantidad_kg_unid', 'Precio_Unitario_USD', 'Monto_Venta_USD'))
    flow = ledger.create_sheet('Operacion')
    flow.append(('Fecha', 'SKU', 'Unidad', 'Stock_inicial_PT', 'Fabricado', 'Entrega_bruta', 'Devuelto_cliente',
        'Merma_PT', 'Stock_final_PT', 'MP_recibida_kg', 'MP_devuelta_proveedor_kg', 'MP_consumida_kg',
        'Empaques_recibidos_unid', 'Empaques_consumidos_unid', 'Factura_cliente_USD', 'Credito_cliente_USD',
        'Cobro_USD', 'Factura_proveedor_USD', 'Credito_proveedor_USD', 'Pago_proveedor_USD', 'Diferencia_stock'))
    count = 0
    total_sales, total_credits, total_cash = Decimal(0), Decimal(0), Decimal(0)
    for offset in range((last - first).days + 1):
        day = first + timedelta(days=offset)
        for pi, product in enumerate(PRODUCTS):
            delivered = 0
            for ci, client in enumerate(CLIENTS):
                qty = 1 + ((offset + ci + pi) % 5)
                price = Decimal(str(product['price']))
                sheet.append((day, client, f'Punto LAB-{ci % 3 + 1:03d} (simulado)',
                    'Distribución LAB simulada', 'Venta LAB simulada', product['name'],
                    'Avícola LAB simulada', qty, float(price), float(qty * price)))
                delivered += qty
                count += 1
            returned, scrap = (1, 1) if day.weekday() == 6 else (0, 0)
            produced = delivered - returned + scrap
            stock = 20 + produced + returned - delivered - scrap
            consumed = Decimal(produced) * Decimal('1.15')
            supplier_return = Decimal('1') if day.day == 1 else Decimal(0)
            received = consumed + supplier_return
            packs = produced * product['pack_per_unit']
            invoice = Decimal(delivered) * price
            credit = Decimal(returned) * price
            supplier_invoice = received * Decimal('1.5') + Decimal(packs) * Decimal('.2')
            supplier_credit = supplier_return * Decimal('1.5')
            assert stock == 20 and received - supplier_return - consumed == 0
            assert invoice >= credit and supplier_invoice >= supplier_credit
            flow.append((day, product['sku'], product['unit'], 20, produced, delivered, returned, scrap, stock,
                float(received), float(supplier_return), float(consumed), packs, packs, float(invoice),
                float(credit), float(invoice - credit), float(supplier_invoice), float(supplier_credit),
                float(supplier_invoice - supplier_credit), 0))
            total_sales += invoice
            total_credits += credit
            total_cash += invoice - credit
    assert total_sales - total_credits == total_cash
    for book in (sales, ledger):
        origin = book.create_sheet('Origen_simulado')
        origin.append(('Origen', 'SIMULADO: generado para pruebas, no movimientos históricos reales ni importación Odoo'))
        origin.append(('Periodo', f'{first} a {last}: 18 meses calendario'))
        origin.append(('Coherencia', 'PT: inicial + fabricado + devuelto - entregado - merma = final; MP: recibido - devolución - consumo = 0'))
        origin.append(('Finanzas', 'Cobro = factura - crédito; pago proveedor = factura proveedor - crédito proveedor; sin impuestos'))
        origin.append(('Stock', '20 por producto es referencia sintética de planificación, no stock ni reservas actuales de Odoo'))
        origin.append(('Reglas', '1 a 5 por cliente/producto/día; devolución y merma 1 los domingos; MP1.15kg y empaque1 por PT nuevo'))
    sales.save(sales_path)
    ledger.save(ledger_path)
    stock_book = Workbook()
    stock = stock_book.active
    stock.title = 'Inventario'
    stock.append(('Cliente', 'Producto', 'Existencia_Disponible', 'Fecha_Corte', 'Unidad'))
    for ci, client in enumerate(CLIENTS):
        for product in PRODUCTS:
            # Allocated planning shares total 20 per SKU; they are not ten copies of shared physical stock.
            stock.append((client, product['name'], 2, last, product['unit']))
    origin = stock_book.create_sheet('Origen_simulado')
    origin.append(('Origen', 'SIMULADO: cuotas de planificación de 2 por cliente, total20porSKU; Odoo mantiene stock y reservas'))
    stock_book.save(inventory_path)
    stock_book.close()
    admin = get_user_model().objects.get(username='lab_admin_01', is_staff=True)
    existing_accounts = list(get_user_model().objects.values_list('pk', 'username', 'password'))
    with transaction.atomic():
        dataset = import_sales(sales_path.read_bytes(), sales_path.name, admin, verified=True, full_snapshot_confirmed=True)
        inventory = import_inventory(inventory_path.read_bytes(), inventory_path.name, admin, unit_match_confirmed=True)
        dataset.notes = 'SIMULADO. 10 clientes, 8 productos y operación diaria reconciliada; no información empresarial.'
        dataset.save(update_fields=['notes'])
        for username, clients in PORTFOLIOS.items():
            account = get_user_model().objects.get(username=username)
            for client in clients:
                assignment, created = Assignment.objects.get_or_create(client=client, defaults={'user': account})
                assert assignment.user_id == account.id, 'Preserve existing portfolios; reconcile unexpected assignment first'
        for n, client in enumerate(CLIENTS, 1):
            mapping, created = OdooCustomerMapping.objects.get_or_create(client=client, defaults={'code': f'VIT-LAB-CLIENT-{n:03d}'})
            assert mapping.code == f'VIT-LAB-CLIENT-{n:03d}'
        for product in PRODUCTS:
            mapping, created = OdooProductMapping.objects.get_or_create(product=product['name'], defaults={
                'sku': product['sku'], 'unit': product['uom'], 'source_unit': product['unit']})
            assert (mapping.sku, mapping.unit, mapping.source_unit) == (product['sku'], product['uom'], product['unit'])
    assert existing_accounts == list(get_user_model().objects.values_list('pk', 'username', 'password'))
    assert dataset.row_count == count and Assignment.objects.count() == 10
    previous = {k: info[k] for k in ('row_count', 'sales_dataset_id', 'inventory_batch_id')}
    previous['sources'] = list(info['sources'])
    info.update(version=2, row_count=count, sales_dataset_id=dataset.id, inventory_batch_id=inventory.id,
        portfolios=PORTFOLIOS, products=PRODUCTS, previous_snapshot=previous)
    info['sources'] += [{'path': str(path), 'sha256': sha256(path.read_bytes()).hexdigest()}
                        for path in (sales_path, inventory_path, ledger_path)]
    marker.write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding='utf-8')
    evidence = {'status': 'passed', 'sales_rows': count, 'operation_rows': ((last - first).days + 1) * 8,
        'customers': 10, 'finished_products': 8, 'calendar_months': 18, 'first_date': str(first), 'last_date': str(last),
        'invoice_total_usd': str(total_sales), 'credit_total_usd': str(total_credits), 'collections_total_usd': str(total_cash),
        'stock_flow_equations': True, 'customer_and_supplier_cash_equations': True,
        'original_sources_and_accounts_preserved': True,
        'notice': 'Historical ledger is synthetic and internally coherent; native Odoo operation is tested separately.'}
    (ROOT / '.local-odoo/evidence/expanded-history.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'PASS: {count} synthetic sales rows, 10 customers, 8 products; stock and cash flows reconcile')


def import_native_inventory():
    from preparar_smartorder_lab import LAB_HOME, bootstrap
    bootstrap()
    from django.contrib.auth import get_user_model
    from operations.models import InventoryBatch
    from operations.services import import_inventory

    proof = json.loads((ROOT / '.local-odoo/evidence/inventory-export.json').read_text(encoding='utf-8'))
    assert proof['status'] == 'passed' and proof['read_only_transaction'] and proof['database'] == 'vitali_lab'
    path = Path(proof['path']).resolve()
    assert path.is_relative_to((LAB_HOME / 'fuentes_simuladas').resolve()) and 'SIMULADO' in path.name
    raw = path.read_bytes()
    assert sha256(raw).hexdigest() == proof['source_sha256']
    marker = LAB_HOME / 'laboratorio.json'
    info = json.loads(marker.read_text(encoding='utf-8'))
    existing = InventoryBatch.objects.filter(digest=proof['source_sha256']).first()
    if existing:
        print('PRESERVED: reviewed native inventory snapshot already imported')
        return
    admin = get_user_model().objects.get(username='lab_admin_01', is_staff=True)
    inventory = import_inventory(raw, path.name, admin, unit_match_confirmed=True)
    assert inventory.items.count() == proof['stock_rows'] == 80
    inventory.notes += '\nSIMULADO ODOO: cuotas del stock nativo disponible por punto, sin duplicar reservas; revisar nuevos snapshots antes de importar.'
    inventory.save(update_fields=['notes'])
    info.update(inventory_batch_id=inventory.pk, native_inventory_snapshot=proof,
        inventory_notice='Reviewed Odoo snapshot: available quantities allocated by sales point. It is dated planning information, not a reservation.')
    info['sources'].append({'path': str(path), 'sha256': proof['source_sha256']})
    marker.write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding='utf-8')
    print('PASS: reviewed native Odoo inventory imported through the shared administrative service; 80 pairs, both interfaces')


if 'env' in globals():
    prepare_odoo(env)
elif __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inventario-odoo', action='store_true', help='Import the validated native inventory export into the separate lab')
    args = parser.parse_args()
    import_native_inventory() if args.inventario_odoo else prepare_smartorder()
