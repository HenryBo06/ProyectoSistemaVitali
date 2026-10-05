"""Prepare a separate, synthetic SmartOrder database; never import business data."""
from datetime import date, timedelta
from hashlib import sha256
import json
import os
from pathlib import Path
import secrets
import sys

ROOT = Path(__file__).resolve().parents[1]
LAB_HOME = ROOT / ".local-web-lab"
LAB_DB = LAB_HOME / "smartorder.sqlite3"
CLIENTS = tuple(f"Cliente Vitali LAB-{index:03d} (simulado)" for index in range(1, 6))
PRODUCT = "Producto avícola LAB (simulado)"
ACCOUNTS = ("lab_admin_01", "lab_admin_02", "lab_vendedor_01", "lab_vendedor_02", "lab_vendedor_03")
PORTFOLIOS = {"lab_vendedor_01": (CLIENTS[0], CLIENTS[3]),
              "lab_vendedor_02": (CLIENTS[1], CLIENTS[4]), "lab_vendedor_03": (CLIENTS[2],)}

# ponytail: the persisted catalog governs checks after expansion; no second user model.
prepared_marker = LAB_HOME / "laboratorio.json"
if prepared_marker.is_file():
    prepared = json.loads(prepared_marker.read_text(encoding="utf-8"))
    if prepared.get("kind") == "SmartOrder synthetic laboratory" and prepared.get("version") == 2:
        PORTFOLIOS = {name: tuple(clients) for name, clients in prepared["portfolios"].items()}
        CLIENTS = tuple(sorted({client for clients in PORTFOLIOS.values() for client in clients}))


def bootstrap(*, odoo=False):
    """Fail closed when a caller supplies a different database or upload directory."""
    for name, expected in (("SMARTORDER_WEB_HOME", LAB_HOME), ("SMARTORDER_DB", LAB_DB)):
        supplied = os.environ.get(name)
        if supplied and Path(supplied).resolve() != expected.resolve():
            raise RuntimeError(f"{name} must point to the separate .local-web-lab laboratory")
        os.environ[name] = str(expected)
    os.environ["SMARTORDER_ODOO_CONFIG"] = str(ROOT / ".local-odoo" / "smartorder.json") if odoo else ""
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "vitali_web.settings")
    # Synthetic checks must not send messages through inherited notification credentials.
    os.environ.pop("SMARTORDER_TELEGRAM_BOT_TOKEN", None)
    os.environ.pop("SMARTORDER_TELEGRAM_CHAT_ID", None)
    sys.path.insert(0, str(ROOT))
    import django
    django.setup()
    from django.conf import settings
    if (Path(settings.DATABASES["default"]["NAME"]).resolve() != LAB_DB.resolve()
            or Path(settings.MEDIA_ROOT).resolve() != (LAB_HOME / "uploads").resolve()):
        raise RuntimeError("Django is already configured for a different database or media directory")
    from django.core.management import call_command
    call_command("migrate", interactive=False, verbosity=0)


def prepare():
    bootstrap()
    from django.contrib.auth import get_user_model
    from django.db import transaction
    from openpyxl import Workbook
    from operations.models import Assignment, OdooCustomerMapping, OdooProductMapping, SalesDataset
    from operations.services import import_inventory, import_sales

    marker = LAB_HOME / "laboratorio.json"
    if marker.is_file():
        info = json.loads(marker.read_text(encoding="utf-8"))
        if info.get("kind") != "SmartOrder synthetic laboratory" or info.get("database") != str(LAB_DB):
            raise RuntimeError("Unrecognized laboratory marker; nothing was replaced")
        if not SalesDataset.objects.filter(pk=info["sales_dataset_id"]).exists():
            raise RuntimeError("The prepared dataset is missing; inspect the laboratory before recreating it")
        print("PRESERVED: isolated synthetic laboratory already prepared")
        return info
    User = get_user_model()
    if User.objects.exists() or SalesDataset.objects.exists():
        raise RuntimeError("Unrecognized existing laboratory data; nothing was replaced")

    sources = LAB_HOME / "fuentes_simuladas"
    sources.mkdir(parents=True, exist_ok=True)
    credential_path = LAB_HOME / "credentials.json"
    if credential_path.exists():
        credentials = json.loads(credential_path.read_text(encoding="utf-8"))
        if set(credentials) != set(ACCOUNTS):
            raise RuntimeError("Unrecognized credential file; nothing was replaced")
    else:
        credentials = {name: secrets.token_urlsafe(24) for name in ACCOUNTS}
        with credential_path.open("x", encoding="utf-8") as output:
            json.dump(credentials, output, indent=2)

    today = date.today()
    month_index = today.year * 12 + today.month - 1 - 17
    first = date(month_index // 12, month_index % 12 + 1, 1)
    book = Workbook()
    sheet = book.active
    sheet.title = "Ventas"
    sheet.append(("Fecha", "Cliente", "Zona_Geografica", "Canal_Distribucion", "Canal_Venta",
                  "Producto", "Categoria", "Cantidad_kg_unid", "Precio_Unitario_USD", "Monto_Venta_USD"))
    for offset in range((today - first).days + 1):
        day = first + timedelta(days=offset)
        for index, client in enumerate(CLIENTS):
            quantity = 1 + ((offset + index) % 5)
            sheet.append((day, client, "Zona LAB simulada", "Distribución LAB simulada", "Venta LAB simulada",
                          PRODUCT, "Avícola LAB simulada", quantity, 4, quantity * 4))
    origin = book.create_sheet("Origen_simulado")
    origin.append(("Origen", "Generado exclusivamente para pruebas; sin filas empresariales"))
    origin.append(("Regla", "1 + ((día desde inicio + índice cliente) % 5) kg; precio USD 4; cinco clientes"))
    origin.append(("Período", f"{first.isoformat()} a {today.isoformat()}: 18 meses calendario incluyendo el actual"))
    sales_path = sources / "ventas_SIMULADAS_18_meses.xlsx"
    book.save(sales_path)
    book.close()

    book = Workbook()
    stock = book.active
    stock.title = "Inventario"
    stock.append(("Cliente", "Producto", "Existencia_Disponible", "Fecha_Corte", "Unidad"))
    for index, client in enumerate(CLIENTS):
        stock.append((client, PRODUCT, 5 if index == 0 else 0, today, "kg"))
    origin = book.create_sheet("Origen_simulado")
    origin.append(("Origen", "Snapshot sintético de planificación; no copia de existencias reales"))
    origin.append(("Reserva", "Odoo mantiene las reservas físicas compartidas; este archivo no reserva existencias"))
    inventory_path = sources / "inventario_SIMULADO.xlsx"
    book.save(inventory_path)
    book.close()

    sales_raw, inventory_raw = sales_path.read_bytes(), inventory_path.read_bytes()
    with transaction.atomic():
        accounts = {name: User.objects.create_user(name, password=credentials[name],
                    is_staff=name.startswith("lab_admin_")) for name in ACCOUNTS}
        admin = accounts[ACCOUNTS[0]]
        dataset = import_sales(sales_raw, sales_path.name, admin, verified=True, full_snapshot_confirmed=True)
        inventory = import_inventory(inventory_raw, inventory_path.name, admin, unit_match_confirmed=True)
        dataset.notes = "SIMULADO. Generación determinista para pruebas; ver hoja Origen_simulado y laboratorio.json."
        dataset.save(update_fields=["notes"])
        inventory.notes += "\nSIMULADO. Snapshot de planificación; Odoo gobierna reservas físicas."
        inventory.save(update_fields=["notes"])
        for name, clients in PORTFOLIOS.items():
            Assignment.objects.bulk_create([Assignment(user=accounts[name], client=client) for client in clients])
        OdooProductMapping.objects.create(product=PRODUCT, sku="VIT-LAB-PT-001",
            unit="uom.product_uom_kgm", source_unit="kg")
        OdooCustomerMapping.objects.bulk_create([OdooCustomerMapping(client=client,
            code=f"VIT-LAB-CLIENT-{index:03d}") for index, client in enumerate(CLIENTS, 1)])
    info = {"kind": "SmartOrder synthetic laboratory", "version": 1, "database": str(LAB_DB),
        "uploads": str(LAB_HOME / "uploads"), "first_date": first.isoformat(), "last_date": today.isoformat(),
        "calendar_months": 18, "row_count": dataset.row_count, "sales_dataset_id": dataset.pk,
        "inventory_batch_id": inventory.pk, "accounts": list(ACCOUNTS), "portfolios": PORTFOLIOS,
        "products": [{"name": PRODUCT, "sku": "VIT-LAB-PT-001", "unit": "kg"}],
        "sources": [{"path": str(path), "sha256": sha256(raw).hexdigest()}
                    for path, raw in ((sales_path, sales_raw), (inventory_path, inventory_raw))],
        "notice": "All rows and accounts are synthetic. Source verification applies only to this laboratory.",
        "inventory_notice": "Planning snapshot only; Odoo is the authority for physical shared stock and reservations."}
    marker.write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"PASS: isolated laboratory; 5 accounts, 5 exclusive client assignments, {dataset.row_count} synthetic sales rows, 18 calendar months")
    print("Credentials and synthetic source files remain under .local-web-lab/; business accounts and files were not imported")
    return info


if __name__ == "__main__":
    prepare()
