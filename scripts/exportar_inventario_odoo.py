"""Odoo shell: reviewable native inventory snapshot; database strictly read-only."""
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal, ROUND_DOWN
from hashlib import sha256
import json
from math import isfinite
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from odoo import fields

assert env.cr.dbname == "vitali_lab", "Synthetic laboratory only"
env.cr.rollback()
env.invalidate_all()
env.cr.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
env.cr.execute("SHOW transaction_read_only")
assert env.cr.fetchone()[0] == "on"
root = Path.cwd()
sys.path.insert(0, str(root))
from smartorder.inventory import load_inventory

now = datetime.now(timezone.utc)
local_clock = env["res.partner"].with_context(tz="America/El_Salvador")
local_today = fields.Datetime.context_timestamp(local_clock, now.replace(tzinfo=None)).date()
company = env.company
warehouse = env.ref("stock.warehouse0")
kg, units = env.ref("uom.product_uom_kgm"), env.ref("uom.product_uom_unit")
catalog = json.loads((root / ".local-odoo/evidence/expanded-catalog.json").read_text(encoding="utf-8"))
assert catalog["status"] == "passed" and catalog["customers"] == 10 and catalog["finished_products"] == 8
points = env["stock.location"].browse(catalog["sales_point_location_ids"]).exists()
transit = env["stock.location"].browse(catalog["transit_location_id"]).exists()
assert len(points) == 3 and len(transit) == 1
for n, point in enumerate(points, 1):
    assert (point.name, point.usage, point.company_id, point.location_id) == (
        f"Punto de venta LAB-{n:03d} (simulado)", "internal", company, warehouse.view_location_id)
assert transit.usage == "transit" and transit.company_id == company
location_roots = [warehouse.lot_stock_id, *points, transit]
scoped_locations = env["stock.location"].search([("id", "child_of", [p.id for p in location_roots]),
    ("company_id", "=", company.id)])
point_children = {p.id: set(env["stock.location"].search([("id", "child_of", p.id)]).ids) for p in points}
customers = []
by_point = defaultdict(list)
for n in range(1, 11):
    ref = f"VIT-LAB-CLIENT-{n:03d}"
    customer = env["res.partner"].search([("ref", "=", ref)])
    assert len(customer) == 1 and customer.name == f"Cliente Vitali LAB-{n:03d} (simulado)"
    point = points[(n - 1) % 3]
    customers.append(customer)
    by_point[point.id].append(customer)
products = []
for n in range(1, 9):
    sku = f"VIT-LAB-PT-{n:03d}"
    prod = env["product.product"].search([("default_code", "=", sku)])
    assert len(prod) == 1 and prod.is_storable and prod.uom_id == (kg if n <= 6 else units)
    products.append(prod)
assert len({p.name for p in products}) == 8
business_models = ("res.partner", "product.product", "stock.location", "stock.quant", "stock.move",
    "stock.picking", "sale.order", "purchase.order", "mrp.production", "account.move", "vitali.integration")
counts_before = {model: env[model].search_count([]) for model in business_models}
original_files = {str(p.relative_to(root)): sha256(p.read_bytes()).hexdigest()
    for p in (root / ".local-web-lab/fuentes_simuladas").glob("*.xlsx")}
original = env["vitali.integration"].search([("sale_id.name", "=", "S00011")])
assert len(original) == 1
original_snapshot = original._snapshot()
original_snapshot.pop("synced_at", None)
original_digest = sha256(json.dumps(original_snapshot, sort_keys=True).encode()).hexdigest()
seed = env["ir.config_parameter"].sudo()
seed_sale = env["sale.order"].browse(int(seed.get_param("vitali_lab.sale_id")))
seed_mo = env["mrp.production"].browse(int(seed.get_param("vitali_lab.manufacturing_id")))
assert seed_sale.state == "sale" and seed_mo.state == "draft"


def number(value):
    assert isfinite(value), value
    return Decimal(str(round(value, 9)))


def allocate(total, owners, prod):
    """Equal planning shares with deterministic remainder; never multiply shared stock."""
    # ponytail: explicit equal shares are planning quotas, not native reservations.
    # Replace this rule with approved business allocations before a real pilot.
    total = max(Decimal(0), number(float(total)))
    assert total == total.quantize(Decimal("0.001")), "Native quantity exceeds SmartOrder's three decimal places"
    step = Decimal("1") if prod.uom_id == units else Decimal("0.001")
    share = (total / len(owners) / step).to_integral_value(rounding=ROUND_DOWN) * step
    portions = [share] * (len(owners) - 1) + [total - share * (len(owners) - 1)]
    if prod.uom_id == units:
        assert all(q == q.to_integral_value() for q in portions), "Fractional native units require review"
    assert sum(portions) == total and min(portions) >= 0
    return list(zip(owners, portions))


def sheet(book, name, headers, rows):
    page = book.create_sheet(name)
    page.append(headers)
    for row in rows:
        page.append(row)
    page.freeze_panes = "A2"
    page.auto_filter.ref = page.dimensions
    for cell in page[1]:
        cell.fill = PatternFill("solid", fgColor="27863F")
        cell.font = Font(color="FFFFFF", bold=True)
    for column in page.columns:
        page.column_dimensions[column[0].column_letter].width = min(65,
            max(17, max(len(str(c.value or "")) for c in column) + 2))
    return page


# Check nonzero shares even when an empty point has no stock or future documents.
kg_example = allocate(Decimal("1"), by_point[points[1].id], products[0])
assert [amount for _, amount in kg_example] == [Decimal(".333"), Decimal(".333"), Decimal(".334")]
unit_example = allocate(Decimal("7"), by_point[points[1].id], products[6])
assert [amount for _, amount in unit_example] == [Decimal("2"), Decimal("2"), Decimal("3")]
assert max(Decimal(0), Decimal("8") - Decimal("5")) == Decimal("3")
assert fields.Datetime.context_timestamp(local_clock, datetime(2026, 10, 4, 3)).date().isoformat() == "2026-10-03"


stocks, allocations, native, flows, move_trace = [], [], [], [], []
stock_balances = []
for location in location_roots:
    for prod in products:
        quants = env["stock.quant"].search([("product_id", "=", prod.id),
            ("company_id", "=", company.id), ("location_id", "child_of", location.id)])
        physical = number(sum(quants.mapped("quantity")))
        reserved = number(sum(quants.mapped("reserved_quantity")))
        available = max(Decimal(0), number(env["stock.quant"]._get_available_quantity(prod, location)))
        unit_name = "kg" if prod.uom_id == kg else "unid"
        native.append((location.id, location.complete_name, location.usage, prod.id, prod.default_code,
            prod.name, unit_name, float(physical), float(reserved), float(physical - reserved),
            float(available), now.isoformat(), ",".join(map(str, quants.ids))))
        if location in points:
            portions = allocate(available, by_point[location.id], prod)
            for customer, amount in portions:
                stocks.append((customer.name, prod.name, float(amount), local_today, unit_name))
                allocations.append((customer.id, customer.ref, customer.name, location.id,
                    location.complete_name, prod.id, prod.default_code, prod.name, prod.uom_id.id,
                    "uom.product_uom_kgm" if prod.uom_id == kg else "uom.product_uom_unit",
                    unit_name, float(amount), float(available), len(portions)))
            assert sum(amount for _, amount in portions) == available
            stock_balances.append({"location_id": location.id, "sku": prod.default_code,
                "unit": unit_name, "available": float(available), "allocated": float(available)})

pending = env["stock.move"].search([("company_id", "=", company.id),
    ("product_id", "in", [p.id for p in products]),
    ("state", "in", ["waiting", "confirmed", "partially_available", "assigned"]),
    "|", ("location_id", "in", scoped_locations.ids), ("location_dest_id", "in", scoped_locations.ids)],
    order="date,id")
future_totals = defaultdict(Decimal)
for move in pending:
    assert move.date, "Native planning date is required; no invented date"
    planned = move.date.replace(tzinfo=timezone.utc)
    due = fields.Datetime.context_timestamp(local_clock, move.date).date()
    prod = move.product_id
    demand = number(move.product_uom._compute_quantity(move.product_uom_qty, prod.uom_id, round=False))
    assigned = number(move.product_uom._compute_quantity(move.quantity, prod.uom_id, round=False))
    unreserved = max(Decimal(0), demand - assigned)
    for point in points:
        is_incoming = move.location_dest_id.id in point_children[point.id] and move.location_id.id not in point_children[point.id]
        is_outgoing = move.location_id.id in point_children[point.id] and move.location_dest_id.id not in point_children[point.id]
        if not is_incoming and not is_outgoing:
            continue
        kind = "Entrada" if is_incoming else "Compromiso_Cliente"
        amount = demand if is_incoming else unreserved
        included = due >= local_today and amount > 0
        unit_name = "kg" if prod.uom_id == kg else "unid"
        move_trace.append((move.id, move.picking_id.id or None, move.picking_id.name or move.reference,
            move.state, prod.default_code, unit_name, point.id, kind,
            move.location_id.id, move.location_dest_id.id, float(demand), float(assigned), float(unreserved),
            planned.isoformat(), due, float(amount) if included else 0,
            "PROYECTADO: documento pendiente; no reserva garantizada" if included else
            "EXCLUIDO: fecha pasada o salida completamente reservada"))
        if included:
            portions = allocate(amount, by_point[point.id], prod)
            for customer, share in portions:
                if share:
                    flows.append((kind, customer.name, prod.name, float(share), due))
            assert sum(share for _, share in portions) == amount
            future_totals[(kind, point.id, prod.default_code, unit_name)] += amount

assert len(stocks) == 80 and len({(row[0], row[1]) for row in stocks}) == 80
book = Workbook()
book.remove(book.active)
sheet(book, "Inventario", ("Cliente", "Producto", "Existencia_Disponible", "Fecha_Corte", "Unidad"), stocks)
sheet(book, "Movimientos", ("Tipo", "Cliente", "Producto", "Cantidad", "Fecha"), flows)
sheet(book, "Asignacion_revisable", ("ID_Cliente_Odoo", "Referencia_Cliente", "Nombre_Cliente", "ID_Ubicacion",
    "Ubicacion", "ID_Producto_Odoo", "SKU", "Nombre_Producto", "ID_Unidad_Odoo", "Unidad_XMLID",
    "Unidad", "Cuota_cliente", "Disponible_punto", "Clientes_del_punto"), allocations)
sheet(book, "Stock_nativo", ("ID_Ubicacion", "Ubicacion", "Uso", "ID_Producto", "SKU", "Nombre_Producto",
    "Unidad", "Fisico", "Reservado", "Fisico_menos_reservado", "Disponible_nativo_no_negativo",
    "Consulta_UTC", "IDs_Quants"), native)
sheet(book, "Movimientos_nativos", ("ID_Movimiento", "ID_Transferencia", "Referencia", "Estado", "SKU", "Unidad",
    "ID_Punto", "Tipo_proyectado", "ID_Origen", "ID_Destino", "Demanda_pendiente", "Asignado_nativo",
    "Demanda_sin_reserva", "Fecha_nativa_UTC", "Fecha_local", "Cantidad_incluida", "Condicion"), move_trace)
sheet(book, "Origen_y_limites", ("Concepto", "Valor"), [
    ("Origen", "SIMULADO: consulta actual y de solo lectura de Odoo Community vitali_lab"),
    ("Consulta_UTC", now.isoformat()), ("Fecha_corte_local", str(local_today)),
    ("Zona_horaria", "America/El_Salvador"),
    ("Asignacion", "001/004/007/010 a punto001;002/005/008 a punto002;003/006/009 a punto003"),
    ("Cuotas", "Reparto uniforme del disponible por punto; sobrante al ultimo cliente; no duplica existencia"),
    ("Unidades", "kg con 3 decimales; unid enteras; totales siempre separados por unidad"),
    ("Reserva", "Cuotas para planificacion revisadas por administrador; no reservas nativas de cliente"),
    ("Movimientos", "Entradas pendientes; salidas solo por demanda aun no reservada para evitar doble descuento"),
    ("Fechas", "Fecha planeada nativa UTC convertida a fecha local; excluye pendientes con fecha pasada"),
    ("Proyeccion", "Documento pendiente no garantiza llegada ni plazo; revisar incidencias y capacidad"),
    ("Saldo_negativo", "Stock_nativo conserva fisico y neto negativos; importacion usa disponible no negativo"),
    ("Actualizacion", "Nueva version necesaria despues de movimientos; no existe sincronizacion automatica del XLSX"),
    ("Datos", "No es historico de 18 meses ni datos empresariales o fiscales reales"),
])
folder = root / ".local-web-lab/fuentes_simuladas"
output = folder / ("inventario_ODOO_SIMULADO_revisable_" + now.strftime("%Y%m%d-%H%M%S-%f") + ".xlsx")
assert not output.exists(), "Existing snapshots are immutable"
book.save(output)
book.close()
raw = output.read_bytes()
parsed = load_inventory(raw)
assert len(parsed.stocks) == 80 and len(parsed.flows) == len(flows)
assert set(parsed.sheets) == {"Inventario", "Movimientos"}
for expected, actual in zip(stocks, parsed.stocks):
    assert (actual.client, actual.product, actual.available, actual.as_of, actual.unit) == expected
for expected, actual in zip(flows, parsed.flows):
    assert actual.quantity == expected[3] and actual.due_date == expected[4]
readback = load_workbook(output, read_only=True, data_only=True)
assert readback["Stock_nativo"].max_row == len(native) + 1 and readback["Asignacion_revisable"].max_row == 81
assert all(cell.data_type != "f" for page in readback for row in page for cell in row)
readback.close()
assert {model: env[model].search_count([]) for model in business_models} == counts_before
after = original._snapshot()
after.pop("synced_at", None)
assert sha256(json.dumps(after, sort_keys=True).encode()).hexdigest() == original_digest
assert seed_sale.state == "sale" and seed_mo.state == "draft"
assert all(sha256((root / p).read_bytes()).hexdigest() == digest for p, digest in original_files.items())
evidence = {"status": "passed", "database": "vitali_lab", "read_only_transaction": True,
    "checked_at_utc": now.isoformat(), "as_of_local": str(local_today), "timezone": "America/El_Salvador",
    "path": str(output), "source_sha256": parsed.digest, "stock_rows": 80,
    "flow_rows": len(flows), "native_stock_rows": len(native), "native_pending_movements": len(move_trace),
    "nonzero_allocation_and_utc_date_self_checks": True,
    "point_allocations": stock_balances,
    "native_totals_by_unit": [{"unit": unit_name,
        "physical": sum(row[7] for row in native if row[6] == unit_name),
        "reserved": sum(row[8] for row in native if row[6] == unit_name),
        "available": sum(row[10] for row in native if row[6] == unit_name)} for unit_name in ("kg", "unid")],
    "future_totals": [{"kind": key[0], "location_id": key[1], "sku": key[2], "unit": key[3],
        "quantity": float(value)} for key, value in future_totals.items()],
    "preserved_business_counts": counts_before, "preserved_integrated_snapshot_sha256": original_digest,
    "original_source_files_preserved": len(original_files), "compatible_shared_inventory_parser": True,
    "reservation_rule": "Available excludes reserves; outgoing projection includes unreserved demand only",
    "notice": "Synthetic reviewable planning shares; no automatic import or native reservation"}
env.cr.rollback()
(root / ".local-odoo/evidence/inventory-export.json").write_text(
    json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
print("PASS: readonly native inventory snapshot, 80 client/product pairs; flows:", len(flows))
print("EXPORT:", output)
print("SHA256:", parsed.digest)
