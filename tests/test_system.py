from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from smartorder.data import HEADERS, load_sales
from smartorder.forecast import DemandForecaster
from smartorder.orders import OperationalInput, export_approved, suggest_order
from smartorder.storage import Store


def workbook_bytes(records):
    book = Workbook()
    sheet = book.active
    sheet.title = "Ventas2025"
    sheet.append(HEADERS)
    for day, client, product, quantity in records:
        sheet.append((datetime.combine(day, datetime.min.time()), client, "Central",
                      "Cadena Supermercados", "Retail", product, "Pollo",
                      quantity, 2.0, quantity * 2.0))
    output = BytesIO()
    book.save(output)
    return output.getvalue()


class SalesTests(unittest.TestCase):
    def test_duplicate_transactions_and_missing_day(self):
        start = date(2025, 1, 1)
        raw = workbook_bytes([
            (start, "Cliente A", "Pollo", 2),
            (start, "Cliente A", "Pollo", 3),
            (start + timedelta(days=2), "Cliente A", "Pollo", 4),
        ])
        data = load_sales(raw)
        self.assertEqual(len(data.rows), 3)
        self.assertEqual(data.daily()["Cantidad_kg_unid"].tolist(), [5, 0, 4])

    def test_rejects_inconsistent_amount(self):
        raw = bytearray(workbook_bytes([(date(2025, 1, 1), "A", "Pollo", 1)]))
        book = load_workbook(BytesIO(raw))
        book.active["J2"] = 500
        stream = BytesIO()
        book.save(stream)
        with self.assertRaisesRegex(ValueError, "monto no coincide"):
            load_sales(stream.getvalue())

    def test_rejects_corrupt_xlsx(self):
        with self.assertRaisesRegex(ValueError, "XLSX legible"):
            load_sales(b"PKnot-a-real-workbook")


class ForecastTests(unittest.TestCase):
    def test_uses_real_previous_month_when_recent_sales_unavailable(self):
        start = date(2025, 1, 1)
        records = [(start + timedelta(days=i), "Cliente A", "Pollo", 10 if i % 7 == 0 else 0)
                   for i in range(365)]
        data = load_sales(workbook_bytes(records))
        report = DemandForecaster().run(data, date(2026, 9, 25))
        row = report.rows.iloc[0]
        self.assertEqual(row["Metodo"], "Mismo mes del año anterior")
        self.assertAlmostEqual(row["Pronostico_7d"], row["Promedio_semanal_mismo_mes"], places=2)
        self.assertTrue(row["Pronostico_7d"] >= 0)
        self.assertEqual(report.validation_windows, 8)


class OrderTests(unittest.TestCase):
    def test_rounding_and_zero_demand(self):
        values = OperationalInput.create("A", "Pollo", "kg", 20, 10, 5, 12, 6, date.today())
        self.assertEqual(suggest_order(30, values).suggested, Decimal("12"))
        self.assertEqual(suggest_order(1, values).suggested, Decimal("0"))
        with self.assertRaisesRegex(ValueError, "mayor que cero"):
            OperationalInput.create("A", "Pollo", "kg", 0, 0, 0, 0, 0, date.today())

    def test_excel_export_escapes_formula_text(self):
        row = {"fecha_decision": "2026-09-25", "usuario": "vendedor",
               "cliente": "=HYPERLINK(\"bad\")", "producto": "Pollo",
               "unidad": "kg", "fecha_pronostico": "2026-09-25",
               "pronostico": 10, "sugerido": 6, "aprobado": 6, "motivo": ""}
        book = load_workbook(BytesIO(export_approved([row])))
        self.assertEqual(book.active["C2"].data_type, "s")


class AccessTests(unittest.TestCase):
    def test_vendor_cannot_use_unassigned_client_or_admin_actions(self):
        with TemporaryDirectory() as folder:
            store = Store(Path(folder) / "test.sqlite3")
            admin = store.bootstrap_admin("admin", "long-admin-password")
            seller = store.create_user(admin, "seller", "long-seller-password", "vendedor")
            store.set_assignments(admin, seller, ["Cliente A"])
            self.assertEqual(store.authenticate("seller", "long-seller-password")["id"], seller)
            with self.assertRaises(PermissionError):
                store.create_user(seller, "another", "another-password", "admin")
            wrong = OperationalInput.create("Cliente B", "Pollo", "kg", 1, 0, 2, 0, 1, date.today())
            with self.assertRaises(PermissionError):
                store.save_operational(seller, wrong)
            right = OperationalInput.create("Cliente A", "Pollo", "kg", 1, 0, 2, 0, 1, date.today())
            store.save_operational(seller, right)
            with self.assertRaisesRegex(ValueError, "Explique"):
                store.save_decision(seller, dataset_digest="abc", client="Cliente A",
                                    product="Pollo", forecast_date=date.today(), forecast=10,
                                    suggested=11, approved=12, operational=right, reason="")


if __name__ == "__main__":
    unittest.main()
