"""Checks at the workbook and messaging trust boundaries."""

from datetime import date
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch
import os

from openpyxl import Workbook, load_workbook

from smartorder.data import load_sales
from smartorder.notifications import notify_internal
from smartorder.orders import export_production


class BoundaryTests(TestCase):
    def test_customer_type_is_not_mislabeled_as_sales_channel(self):
        book = Workbook()
        sheet = book.active
        sheet.append(("Fecha", "Cliente", "Producto", "Cantidad_kg_unid",
                      "Precio_Unitario_USD", "Tipo_Cliente"))
        sheet.append((date(2025, 1, 1), "A", "Pollo", 2, 3, "Mayorista"))
        output = BytesIO()
        book.save(output)
        row = load_sales(output.getvalue()).rows.iloc[0]
        self.assertEqual(row["Canal_Distribucion"], "Mayorista")
        self.assertEqual(row["Canal_Venta"], "No especificado")
        self.assertEqual(row["Monto_Venta_USD"], 6)

    def test_export_escapes_excel_formula_cells(self):
        raw = export_production([{
            "estado": "Aprobado", "id": 1, "cliente": "A", "producto": "=1+1",
            "unidad": "kg", "cantidad_produccion": 2, "fecha_requerida": "2026-09-28",
            "fecha_pronostico": "2026-09-28", "usuario": "vendedor",
            "revisor": "admin", "fecha_revision": "2026-09-28", "nota_revision": "",
        }])
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        self.assertEqual(book["Pedidos para produccion"]["C2"].value, "'=1+1")
        self.assertEqual(book["Resumen por producto"]["A2"].value, "'=1+1")
        book.close()

    def test_telegram_sends_only_allowed_summary_when_configured(self):
        with patch.dict(os.environ, {
            "SMARTORDER_TELEGRAM_BOT_TOKEN": "test-token",
            "SMARTORDER_TELEGRAM_CHAT_ID": "123",
        }):
            with patch("smartorder.notifications.urlopen", return_value=BytesIO(b'{"ok": true}')) as send:
                result = notify_internal("Pedido pendiente de revisión", "http://127.0.0.1:8000/pedidos/")
            self.assertEqual(result, (True, "Aviso enviado."))
            self.assertIn(b"Pedido+pendiente", send.call_args.args[0].data)
            with patch("smartorder.notifications.urlopen") as forbidden_send:
                self.assertFalse(notify_internal("Enviar datos privados", "http://localhost/")[0])
                forbidden_send.assert_not_called()
