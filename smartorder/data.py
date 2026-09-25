"""Lectura estricta del histórico comercial recibido en Excel."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


HEADERS = (
    "Fecha", "Cliente", "Zona_Geografica", "Canal_Distribucion", "Canal_Venta",
    "Producto", "Categoria", "Cantidad_kg_unid", "Precio_Unitario_USD",
    "Monto_Venta_USD",
)
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 150_000
MAX_DAILY_CELLS = 250_000


@dataclass(frozen=True)
class SalesData:
    rows: pd.DataFrame
    digest: str
    sheets: tuple[str, ...]
    recalculated_amounts: int

    @property
    def first_date(self) -> date:
        return self.rows["Fecha"].min().date()

    @property
    def last_date(self) -> date:
        return self.rows["Fecha"].max().date()

    @property
    def clients(self) -> list[str]:
        return sorted(self.rows["Cliente"].unique().tolist())

    def daily(self) -> pd.DataFrame:
        """Una fila por día, cliente y producto; ausencia significa cero ventas."""
        pairs = self.rows[["Cliente", "Producto"]].drop_duplicates()
        days = pd.DataFrame({"Fecha": pd.date_range(self.first_date, self.last_date)})
        grid = pairs.merge(days, how="cross")
        totals = self.rows.groupby(["Fecha", "Cliente", "Producto"], as_index=False)[
            "Cantidad_kg_unid"
        ].sum()
        return grid.merge(totals, on=["Fecha", "Cliente", "Producto"], how="left").fillna(
            {"Cantidad_kg_unid": 0.0}
        ).sort_values(["Cliente", "Producto", "Fecha"]).reset_index(drop=True)


def _as_date(value: object, row: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError as exc:
            raise ValueError(f"Fila {row}: Fecha debe tener formato AAAA-MM-DD.") from exc
    raise ValueError(f"Fila {row}: Fecha inválida.")


def _number(value: object, name: str, row: int) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Fila {row}: {name} debe ser numérico.")
    result = float(value)
    if not 0 <= result < float("inf"):
        raise ValueError(f"Fila {row}: {name} debe ser finito y no negativo.")
    return result


def load_sales(source: bytes | str | Path) -> SalesData:
    raw = source if isinstance(source, bytes) else Path(source).read_bytes()
    if len(raw) > MAX_BYTES or not raw.startswith(b"PK"):
        raise ValueError("El archivo debe ser un XLSX válido de 20 MB o menos.")

    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
        raise ValueError("El archivo no contiene un libro XLSX legible.") from exc
    records: list[dict[str, object]] = []
    used_sheets: list[str] = []
    recalculated = 0
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, None)
            if header is None or not set(HEADERS).issubset(header):
                continue
            positions = {name: header.index(name) for name in HEADERS}
            used_sheets.append(sheet.title)
            for excel_row, values in enumerate(iterator, start=2):
                if not any(value is not None for value in values):
                    continue
                if len(records) >= MAX_ROWS:
                    raise ValueError(f"El límite es {MAX_ROWS:,} registros.")
                item = {name: values[index] if index < len(values) else None
                        for name, index in positions.items()}
                for name in ("Cliente", "Zona_Geografica", "Canal_Distribucion",
                             "Canal_Venta", "Producto", "Categoria"):
                    text = item[name]
                    if not isinstance(text, str) or not text.strip():
                        raise ValueError(f"Hoja {sheet.title}, fila {excel_row}: falta {name}.")
                    item[name] = text.strip()
                item["Fecha"] = _as_date(item["Fecha"], excel_row)
                quantity = _number(item["Cantidad_kg_unid"], "Cantidad_kg_unid", excel_row)
                price = _number(item["Precio_Unitario_USD"], "Precio_Unitario_USD", excel_row)
                item["Cantidad_kg_unid"] = quantity
                item["Precio_Unitario_USD"] = price
                calculated_amount = quantity * price
                if item["Monto_Venta_USD"] is None:
                    # ponytail: aceptar fórmulas sin caché y recalcular; validar la fuente
                    # original si en el futuro se necesitan ajustes/descuentos por línea.
                    item["Monto_Venta_USD"] = calculated_amount
                    recalculated += 1
                else:
                    amount = _number(item["Monto_Venta_USD"], "Monto_Venta_USD", excel_row)
                    if abs(amount - calculated_amount) > 0.02:
                        raise ValueError(
                            f"Hoja {sheet.title}, fila {excel_row}: monto no coincide "
                            "con cantidad × precio."
                        )
                    item["Monto_Venta_USD"] = amount
                records.append(item)
    finally:
        workbook.close()
    if not used_sheets:
        raise ValueError("Ninguna hoja contiene las diez columnas esperadas de ventas.")
    if not records:
        raise ValueError("El Excel no contiene ventas.")
    frame = pd.DataFrame.from_records(records, columns=HEADERS)
    frame["Fecha"] = pd.to_datetime(frame["Fecha"])
    pairs = frame[["Cliente", "Producto"]].drop_duplicates().shape[0]
    days = (frame["Fecha"].max() - frame["Fecha"].min()).days + 1
    if pairs * days > MAX_DAILY_CELLS:
        raise ValueError(
            "El histórico supera el límite de 250,000 combinaciones diarias "
            "cliente–producto; revise fechas extremas o divida el análisis."
        )
    return SalesData(frame, sha256(raw).hexdigest(), tuple(used_sheets), recalculated)
