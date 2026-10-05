"""Plantilla y lectura estricta de existencias oficiales por cliente y producto."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from math import isfinite
import re
import unicodedata
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 150_000


@dataclass(frozen=True)
class StockRow:
    client: str
    product: str
    available: float
    as_of: date
    unit: str


@dataclass(frozen=True)
class DatedRow:
    kind: str
    client: str
    product: str
    quantity: float
    due_date: date


@dataclass(frozen=True)
class InventoryData:
    digest: str
    stocks: tuple[StockRow, ...]
    flows: tuple[DatedRow, ...]
    sheets: tuple[str, ...]


def _key(value: object) -> str:
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")


def _columns(row: tuple[object, ...], aliases: dict[str, tuple[str, ...]]) -> dict[str, int]:
    found = {_key(value): pos for pos, value in enumerate(row) if value is not None}
    return {field: next((found[key] for key in names if key in found), -1)
            for field, names in aliases.items()}


def _value(row: tuple[object, ...], column: int) -> object:
    return row[column] if 0 <= column < len(row) else None


def _text(value: object, field: str, line: int) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"Fila {line}: {field} es obligatorio.")
    return value.strip()


def _quantity(value: object, field: str, line: int) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Fila {line}: {field} debe ser un número no negativo.")
    number = float(value)
    if not isfinite(number) or number < 0:
        raise ValueError(f"Fila {line}: {field} debe ser un número finito no negativo.")
    return number


def _date(value: object, field: str, line: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    raise ValueError(f"Fila {line}: {field} debe ser una fecha válida.")


STOCK_COLUMNS = {
    "client": ("cliente", "nombre_cliente"),
    "product": ("producto", "nombre_producto"),
    "available": ("existencia_disponible", "disponible", "inventario_disponible", "stock"),
    "as_of": ("fecha_corte", "fecha_inventario", "fecha"),
    "unit": ("unidad", "unidad_medida"),
}
FLOW_COLUMNS = {
    "kind": ("tipo", "movimiento"),
    "client": ("cliente", "nombre_cliente"),
    "product": ("producto", "nombre_producto"),
    "quantity": ("cantidad", "cantidad_pendiente"),
    "due_date": ("fecha", "fecha_prevista", "fecha_compromiso"),
}
FLOW_TYPES = {
    "entrada": "incoming",
    "ingreso": "incoming",
    "incoming": "incoming",
    "compromiso_cliente": "customer_commitment",
    "pedido_cliente": "customer_commitment",
    "customer_commitment": "customer_commitment",
}


def load_inventory(raw: bytes) -> InventoryData:
    if len(raw) > MAX_BYTES or not raw.startswith(b"PK"):
        raise ValueError("El inventario debe ser un archivo XLSX válido de 20 MB o menos.")
    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
        raise ValueError("El archivo de inventario no es un XLSX legible.") from exc
    stocks: list[StockRow] = []
    flows: list[DatedRow] = []
    sheets: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            stock_cols = _columns(header, STOCK_COLUMNS)
            flow_cols = _columns(header, FLOW_COLUMNS)
            is_stock = all(stock_cols[key] >= 0 for key in ("client", "product", "available", "as_of"))
            is_flow = all(pos >= 0 for pos in flow_cols.values())
            if not is_stock and not is_flow:
                continue
            sheets.append(sheet.title)
            for line, row in enumerate(rows, start=2):
                if not any(value is not None for value in row):
                    continue
                if len(stocks) + len(flows) >= MAX_ROWS:
                    raise ValueError(f"El inventario supera {MAX_ROWS:,} filas.")
                if is_stock:
                    stocks.append(StockRow(
                        _text(_value(row, stock_cols["client"]), "Cliente", line),
                        _text(_value(row, stock_cols["product"]), "Producto", line),
                        _quantity(_value(row, stock_cols["available"]), "Existencia disponible", line),
                        _date(_value(row, stock_cols["as_of"]), "Fecha de corte", line),
                        str(_value(row, stock_cols["unit"]) or "").strip(),
                    ))
                else:
                    kind = FLOW_TYPES.get(_key(_value(row, flow_cols["kind"])))
                    if kind is None:
                        raise ValueError(f"Fila {line}: tipo de movimiento desconocido.")
                    flows.append(DatedRow(
                        kind,
                        _text(_value(row, flow_cols["client"]), "Cliente", line),
                        _text(_value(row, flow_cols["product"]), "Producto", line),
                        _quantity(_value(row, flow_cols["quantity"]), "Cantidad", line),
                        _date(_value(row, flow_cols["due_date"]), "Fecha", line),
                    ))
    finally:
        workbook.close()
    if any(row.as_of > date.today() for row in stocks):
        raise ValueError("La fecha de corte del inventario no puede estar en el futuro.")
    if not stocks:
        raise ValueError("Falta una hoja de inventario con Cliente, Producto, Existencia_Disponible y Fecha_Corte.")
    duplicates = [(row.client, row.product) for row in stocks]
    if len(duplicates) != len(set(duplicates)):
        raise ValueError("Hay pares Cliente y Producto repetidos en inventario; deje una existencia oficial por par.")
    stock_pairs = set(duplicates)
    if any((row.client, row.product) not in stock_pairs for row in flows):
        raise ValueError("Un movimiento no coincide con ningún Cliente y Producto del inventario.")
    return InventoryData(sha256(raw).hexdigest(), tuple(stocks), tuple(flows), tuple(sheets))


def inventory_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventario"
    sheet.append(("Cliente", "Producto", "Existencia_Disponible", "Fecha_Corte", "Unidad"))
    sheet.freeze_panes = "A2"
    flow = workbook.create_sheet("Movimientos")
    flow.append(("Tipo", "Cliente", "Producto", "Cantidad", "Fecha"))
    flow.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
