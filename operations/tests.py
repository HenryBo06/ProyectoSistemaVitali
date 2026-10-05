"""Integration checks for the local sales-to-production workflow."""

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .models import (
    Assignment, DatedFlow, InventoryAdjustment, InventoryBatch, InventoryItem,
    OrderCoverage, OrderLine, OrderRequest, SalesDataset, SalesLine,
    StockCorrectionRequest,
)
from .services import import_inventory, import_sales


def sales_book(first: date, days: int, *, multiplier: int = 1,
               products: tuple[str, ...] = ("Alitas", "Chorizo")) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Ventas"
    sheet.append(("Fecha", "Cliente", "Producto", "Cantidad_kg_unid",
                  "Precio_Unitario_USD", "Monto_Venta_USD"))
    for offset in range(days):
        day = first + timedelta(days=offset)
        for index, product in enumerate(products, start=1):
            qty = index * multiplier
            sheet.append((day, "Cliente A", product, qty, 10, qty * 10))
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def inventory_book(cutoff: date, *, available: int = 2) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Inventario"
    sheet.append(("Cliente", "Producto", "Existencia_Disponible", "Fecha_Corte", "Unidad"))
    sheet.append(("Cliente A", "Alitas", available, cutoff, "kg"))
    sheet.append(("Cliente A", "Chorizo", available, cutoff, "kg"))
    flow = book.create_sheet("Movimientos")
    flow.append(("Tipo", "Cliente", "Producto", "Cantidad", "Fecha"))
    flow.append(("Entrada", "Cliente A", "Alitas", 1, cutoff + timedelta(days=1)))
    flow.append(("Compromiso_Cliente", "Cliente A", "Alitas", 3,
                 cutoff + timedelta(days=2)))
    output = BytesIO()
    book.save(output)
    return output.getvalue()


class LocalCase:
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.media_override = override_settings(MEDIA_ROOT=self.tmp.name)
        self.media_override.enable()
        self.addCleanup(self.media_override.disable)
        User = get_user_model()
        self.admin = User.objects.create_user(
            "admin_vitali", password="ClaveAdministrativa123", is_staff=True
        )
        self.seller = User.objects.create_user(
            "vendedor_vitali", password="ClaveDelVendedor123"
        )


class DataAndSchemaTests(LocalCase, TestCase):
    def test_sales_and_inventory_etl_preserve_source_and_dated_flows(self):
        today = timezone.localdate()
        sales = import_sales(sales_book(today - timedelta(days=364), 365),
                             "ventas.xlsx", self.admin, verified=True)
        sales.full_snapshot_confirmed = True
        sales.save(update_fields=["full_snapshot_confirmed"])
        self.assertTrue(sales.file.name.startswith("sales/"))
        self.assertEqual(sales.lines.count(), 730)
        self.assertEqual(SalesLine.objects.filter(dataset=sales, product="Alitas").count(), 365)
        stock = import_inventory(inventory_book(today), "inventario.xlsx", self.admin)
        stock.unit_match_confirmed = True
        stock.save(update_fields=["unit_match_confirmed"])
        self.assertTrue(stock.file.name.startswith("inventory/"))
        self.assertEqual(InventoryItem.objects.filter(batch=stock).count(), 2)
        self.assertEqual(
            list(DatedFlow.objects.filter(batch=stock).values_list("kind", flat=True)),
            ["incoming", "customer_commitment"],
        )

    def test_weekly_replacement_preserves_request_pending_revalidation(self):
        first = date(2025, 1, 1)
        old = import_sales(sales_book(first, 40), "primero.xlsx", self.admin, verified=False)
        request = OrderRequest.objects.create(
            seller=self.seller, client="Cliente A", dataset=old,
            required_date=timezone.localdate(),
        )
        OrderLine.objects.create(
            order=request, product="Alitas", method="Pedido manual",
            requested_qty=Decimal("5"), reason="Reposición solicitada",
        )
        new = import_sales(sales_book(first, 40, multiplier=2),
                           "segundo.xlsx", self.admin, verified=False)
        request.refresh_from_db()
        old.refresh_from_db()
        self.assertEqual(request.status, OrderRequest.Status.PENDING)
        self.assertIsNone(request.revalidated_against)
        self.assertFalse(old.active)
        self.assertTrue(new.active)

    def test_only_one_seller_can_own_a_client(self):
        Assignment.objects.create(user=self.seller, client="Cliente A")
        with self.assertRaises(IntegrityError):
            Assignment.objects.create(user=self.admin, client="Cliente A")

    def test_day_level_coverage_is_unique_across_orders(self):
        dataset = import_sales(sales_book(date(2025, 1, 1), 2),
                               "ventas.xlsx", self.admin, verified=False)
        day = timezone.localdate()
        first = OrderRequest.objects.create(
            seller=self.seller, client="Cliente A", dataset=dataset, required_date=day,
        )
        second = OrderRequest.objects.create(
            seller=self.seller, client="Cliente A", dataset=dataset,
            required_date=day + timedelta(days=1),
        )
        line_a = OrderLine.objects.create(
            order=first, product="Alitas", requested_qty=Decimal("5"), method="manual"
        )
        line_b = OrderLine.objects.create(
            order=second, product="Alitas", requested_qty=Decimal("5"), method="manual"
        )
        OrderCoverage.objects.create(line=line_a, client="Cliente A", product="Alitas", day=day)
        with self.assertRaises(IntegrityError):
            OrderCoverage.objects.create(line=line_b, client="Cliente A", product="Alitas", day=day)

    def test_adjustment_preserves_imported_inventory_row(self):
        batch = import_inventory(inventory_book(timezone.localdate()),
                                 "inventario.xlsx", self.admin)
        original = batch.items.get(client="Cliente A", product="Alitas")
        correction = StockCorrectionRequest.objects.create(
            seller=self.seller, client="Cliente A", product="Alitas", unit="kg",
            proposed_available=Decimal("9"), observed_on=timezone.localdate(),
            reason="Conteo físico actualizado",
        )
        InventoryAdjustment.objects.create(
            batch=batch, correction=correction, client="Cliente A", product="Alitas",
            available=Decimal("9"), observed_on=timezone.localdate(), approved_by=self.admin,
        )
        original.refresh_from_db()
        self.assertEqual(original.available, Decimal("2"))
        self.assertEqual(batch.adjustments.get(correction=correction).available, Decimal("9"))


class LocalWorkflowTests(LocalCase, TestCase):
    def _upload_sales(self, raw: bytes, *, verified: bool = False,
                      full_snapshot: bool = False) -> SalesDataset:
        self.client.force_login(self.admin)
        fields = {
            "sales_file": SimpleUploadedFile(
                "ventas.xlsx", raw,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        }
        if verified:
            fields.update(verified="on", confirmed_origin="on")
        if full_snapshot:
            fields["confirmed_full_snapshot"] = "on"
        response = self.client.post(reverse("sales_upload"), fields)
        self.assertEqual(response.status_code, 302)
        return SalesDataset.objects.get(active=True)

    def _upload_inventory(self, raw: bytes, *, units_confirmed: bool = False) -> InventoryBatch:
        self.client.force_login(self.admin)
        fields = {
            "inventory_file": SimpleUploadedFile(
                "inventario.xlsx", raw,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        }
        if units_confirmed:
            fields["unit_match_confirmed"] = "on"
        response = self.client.post(reverse("inventory_upload"), fields)
        self.assertEqual(response.status_code, 302)
        return InventoryBatch.objects.get(active=True)

    def _assign_seller(self):
        self.client.force_login(self.admin)
        response = self.client.post(reverse("users"), {
            "action": "assign", "user_id": self.seller.pk, "clients": ["Cliente A"],
        })
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Assignment.objects.filter(user=self.seller, client="Cliente A").exists())

    @override_settings(DEBUG=True)
    def test_setup_is_local_and_single_use(self):
        get_user_model().objects.all().delete()
        denied = self.client.get(reverse("setup"), REMOTE_ADDR="192.0.2.10", HTTP_HOST="localhost")
        self.assertEqual(denied.status_code, 403)
        response = self.client.post(reverse("setup"), {
            "username": "jefe_vitali", "password": "NuevaClaveLocal#2026",
        }, REMOTE_ADDR="127.0.0.1", HTTP_HOST="localhost")
        self.assertEqual(response.status_code, 302)
        self.assertTrue(get_user_model().objects.get(username="jefe_vitali").is_staff)
        self.assertEqual(self.client.get(reverse("setup"), HTTP_HOST="localhost").status_code, 302)

    @override_settings(SMARTORDER_SETUP_CODE="setup-code-for-tests-only-32-characters")
    def test_remote_setup_requires_the_configured_code(self):
        get_user_model().objects.all().delete()
        remote = {"REMOTE_ADDR": "192.0.2.10", "HTTP_HOST": "localhost"}
        self.assertEqual(self.client.get(reverse("setup"), **remote).status_code, 200)
        with override_settings(SMARTORDER_SETUP_CODE="demasiado-corto"):
            self.assertEqual(self.client.get(reverse("setup"), **remote).status_code, 403)
        invalid = self.client.post(reverse("setup"), {
            "username": "jefe_vitali", "password": "NuevaClaveSegura#2026",
            "setup_code": "incorrecto",
        }, **remote)
        self.assertEqual(invalid.status_code, 200)
        self.assertFalse(get_user_model().objects.exists())
        valid = self.client.post(reverse("setup"), {
            "username": "jefe_vitali", "password": "NuevaClaveSegura#2026",
            "setup_code": "setup-code-for-tests-only-32-characters",
        }, **remote)
        self.assertEqual(valid.status_code, 302)
        self.assertTrue(get_user_model().objects.get(username="jefe_vitali").is_staff)

    def test_seller_cannot_upload_manage_users_review_or_export(self):
        self.client.force_login(self.seller)
        for name in ("sales_upload", "inventory_upload", "inventory_template", "users"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 403, name)
        self.assertEqual(self.client.post(reverse("export_batch"), {"line_ids": ["1"]}).status_code, 403)
        self._upload_sales(sales_book(date(2025, 1, 1), 5))
        self.client.force_login(self.seller)
        self.assertEqual(self.client.get(reverse("order_create") + "?client=Otro").status_code, 403)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("order_create")).status_code, 403)

    def test_post_without_csrf_is_blocked(self):
        strict = Client(enforce_csrf_checks=True)
        strict.force_login(self.admin)
        response = strict.post(reverse("sales_upload"), {"verified": ""})
        self.assertEqual(response.status_code, 403)

    def test_2025_reference_yields_manual_multi_product_review_and_export(self):
        today = date.today()
        year = today.year - 1
        first = date(year, 1, 1)
        days = (date(year + 1, 1, 1) - first).days
        dataset = self._upload_sales(sales_book(first, days))
        self.assertFalse(dataset.verified)
        self._assign_seller()
        self.client.force_login(self.seller)
        recommendation = self.client.get(reverse("recommendations"), {
            "client": "Cliente A", "required_date": today.isoformat(),
        })
        self.assertEqual(recommendation.status_code, 200)
        self.assertTrue(all(card["suggested"] is None
                            for card in recommendation.context["recommendations"]))
        response = self.client.post(reverse("order_create"), {
            "client": "Cliente A", "required_date": today.isoformat(),
            "note": "Revisar para producción",
            "product": ["Alitas", "Chorizo"], "quantity": ["5", "3"],
            "unit": ["kg", "kg"], "reason": ["Pedido manual", "Pedido manual"],
        })
        self.assertEqual(response.status_code, 302)
        order = OrderRequest.objects.get(seller=self.seller)
        self.assertEqual(order.lines.count(), 2)
        self.assertTrue(all(line.suggested_qty is None for line in order.lines.all()))
        lines = {line.product: line for line in order.lines.all()}
        self.client.force_login(self.admin)
        self.client.post(reverse("order_review", args=[order.pk]), {
            "line_id": lines["Alitas"].pk, "decision": "approve",
            "approved_quantity": "4", "review_note": "Ajustado por administración",
        })
        self.client.post(reverse("order_review", args=[order.pk]), {
            "line_id": lines["Chorizo"].pk, "decision": "reject",
            "review_note": "Sin capacidad para este producto",
        })
        order.refresh_from_db()
        self.assertEqual(order.status, OrderRequest.Status.REVIEWED)
        premature = self.client.post(reverse("export_batch"), {
            "line_ids": [str(lines["Alitas"].pk)], "action": "confirm",
        })
        self.assertEqual(premature.status_code, 302)
        lines["Alitas"].refresh_from_db()
        self.assertEqual(lines["Alitas"].status, OrderLine.Status.APPROVED)
        export = self.client.post(reverse("export_batch"), {
            "line_ids": [str(lines["Alitas"].pk)], "action": "download",
        })
        self.assertEqual(export.status_code, 200)
        self.assertTrue(export.content.startswith(b"PK"))
        workbook = load_workbook(BytesIO(export.content), read_only=True)
        self.assertEqual(workbook.active["C2"].value, "Alitas")
        self.assertEqual(workbook.active["E2"].value, 4)
        workbook.close()
        confirmed = self.client.post(reverse("export_batch"), {
            "line_ids": [str(lines["Alitas"].pk)], "action": "confirm",
            "shared_confirm": "1",
        })
        self.assertEqual(confirmed.status_code, 302)
        lines["Alitas"].refresh_from_db()
        order.refresh_from_db()
        self.assertEqual(lines["Alitas"].status, OrderLine.Status.EXPORTED)
        self.assertEqual(order.status, OrderRequest.Status.EXPORTED)

    def test_recent_verified_source_and_stock_can_offer_explained_quantity(self):
        today = date.today()
        dataset = self._upload_sales(sales_book(today - timedelta(days=365), 365),
                                     verified=True, full_snapshot=True)
        self.assertTrue(dataset.verified)
        stock = self._upload_inventory(inventory_book(today), units_confirmed=True)
        self.assertTrue(stock.active)
        self._assign_seller()
        self.client.force_login(self.seller)
        response = self.client.get(reverse("recommendations"), {
            "client": "Cliente A", "required_date": (today + timedelta(days=3)).isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        cards = {card["product"]: card for card in response.context["recommendations"]}
        self.assertContains(response, "decision-suggested")
        self.assertContains(response, "Corte de inventario:")
        self.assertNotContains(response, "Decisión manual")
        self.assertEqual(cards["Alitas"]["state"], "suggested")
        self.assertIsNotNone(cards["Alitas"]["suggested"])
        self.assertNotIn("incoming", cards["Alitas"])
        self.assertNotIn("commitments_before", cards["Alitas"])
        self.assertNotIn("sales_usd_display", cards["Alitas"])
        from .services import recommendation_cards
        internal = next(card for card in recommendation_cards(dataset, today + timedelta(days=3), "Cliente A") if card["product"] == "Alitas")
        self.assertEqual(internal["incoming"], Decimal("1"))
        self.assertEqual(internal["commitments_before"], Decimal("3"))
        self.assertEqual(internal["commitments"], Decimal("0"))
        response = self.client.post(reverse("order_create"), {
            "client": "Cliente A",
            "required_date": (today + timedelta(days=3)).isoformat(),
            "product": ["Alitas", "Chorizo"],
            "quantity": [str(cards["Alitas"]["suggested"]),
                         str(cards["Chorizo"]["suggested"])],
            "unit": ["kg", "kg"], "reason": ["", ""],
        })
        self.assertEqual(response.status_code, 302)
        order = OrderRequest.objects.get(seller=self.seller)
        self.assertEqual(order.lines.count(), 2)
        self.assertTrue(all(line.suggested_qty is not None for line in order.lines.all()))

    def test_unconfirmed_inventory_units_block_automatic_order_quantity(self):
        today = date.today()
        self._upload_sales(sales_book(today - timedelta(days=365), 365,
                                      products=("Alitas",)),
                           verified=True, full_snapshot=True)
        batch = self._upload_inventory(inventory_book(today), units_confirmed=False)
        self.assertFalse(batch.unit_match_confirmed)
        self._assign_seller()
        self.client.force_login(self.seller)
        response = self.client.get(reverse("recommendations"), {
            "client": "Cliente A", "required_date": today.isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        card = response.context["recommendations"][0]
        self.assertIsNone(card["suggested"])
        self.assertIn("unidad", card["inventory_warning"])

    def test_empty_sales_filter_keeps_active_source_and_reset_action(self):
        dataset = self._upload_sales(sales_book(date(2025, 1, 1), 40))
        self.client.force_login(self.admin)
        response = self.client.get(reverse("dashboard"), {
            "start_date": "2024-01-01", "end_date": "2024-01-31",
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["dataset"].pk, dataset.pk)
        self.assertContains(response, "Sin ventas para estos filtros")
        self.assertContains(response, "Restablecer filtros")
        self.assertContains(response, 'id="filter-start"')
        self.assertNotContains(response, "Aún no hay ventas activas")

    def test_role_pages_render_real_metrics_and_seller_products(self):
        today = date.today()
        self._upload_sales(sales_book(today - timedelta(days=365), 365))
        self._upload_inventory(inventory_book(today))
        self._assign_seller()
        self.client.force_login(self.admin)
        for page in ("dashboard", "sales_upload", "inventory_upload", "users",
                     "orders", "correction_requests"):
            response = self.client.get(reverse(page))
            self.assertEqual(response.status_code, 200, page)
        dashboard = self.client.get(reverse("dashboard"))
        self.assertContains(dashboard, "monthlyChart")
        self.assertContains(dashboard, "channelChart")
        self.assertContains(dashboard, "productsChart")
        self.assertContains(dashboard, "Alitas")
        template = self.client.get(reverse("inventory_template"))
        self.assertTrue(template.content.startswith(b"PK"))
        self.client.force_login(self.seller)
        for page in ("dashboard", "recommendations", "order_create",
                     "orders", "correction_requests"):
            response = self.client.get(reverse(page))
            self.assertEqual(response.status_code, 200, page)
        products = self.client.get(reverse("recommendations"))
        self.assertContains(products, "Alitas")
        self.assertNotContains(products, "Referencia histórica")
        self.assertNotContains(products, "WAPE")

    def test_confirmed_unit_cannot_be_changed_in_post(self):
        today = date.today()
        self._upload_sales(sales_book(today - timedelta(days=365), 365),
                           verified=True, full_snapshot=True)
        self._upload_inventory(inventory_book(today), units_confirmed=True)
        self._assign_seller()
        self.client.force_login(self.seller)
        response = self.client.get(reverse("recommendations"), {
            "client": "Cliente A", "required_date": (today + timedelta(days=3)).isoformat(),
        })
        card = {row["product"]: row for row in response.context["recommendations"]}["Alitas"]
        self.assertIsNotNone(card["suggested"])
        created = self.client.post(reverse("order_create"), {
            "client": "Cliente A", "required_date": (today + timedelta(days=3)).isoformat(),
            "product": ["Alitas"], "quantity": [str(card["suggested"])],
            "unit": ["docenas"], "reason": ["manipulado"],
        })
        self.assertEqual(created.status_code, 200)
        self.assertFalse(OrderRequest.objects.exists())

    def test_correction_preserves_import_and_updates_effective_stock(self):
        today = date.today()
        self._upload_sales(sales_book(today - timedelta(days=365), 365))
        batch = self._upload_inventory(inventory_book(today), units_confirmed=True)
        self._assign_seller()
        self.client.force_login(self.seller)
        sent = self.client.post(reverse("correction_requests"), {
            "client": "Cliente A", "product": "Alitas",
            "claimed_stock": "9", "reason": "Conteo verificado hoy",
        })
        self.assertEqual(sent.status_code, 302)
        from .models import StockCorrectionRequest
        correction = StockCorrectionRequest.objects.get()
        self.client.force_login(self.admin)
        accepted = self.client.post(reverse("correction_requests"), {
            "correction_id": correction.pk, "action": "approve",
        })
        self.assertEqual(accepted.status_code, 302)
        self.assertEqual(InventoryItem.objects.get(
            batch=batch, client="Cliente A", product="Alitas"
        ).available, Decimal("2"))
        self.assertEqual(InventoryAdjustment.objects.get(correction=correction).available, Decimal("9"))
        from .services import recommendation_cards
        cards = recommendation_cards(SalesDataset.objects.get(active=True), today, "Cliente A")
        self.assertEqual(next(card for card in cards if card["product"] == "Alitas")["available"],
                         Decimal("9"))

    def test_replaced_sales_file_blocks_review_of_old_request(self):
        first = date(2025, 1, 1)
        self._upload_sales(sales_book(first, 40))
        self._assign_seller()
        self.client.force_login(self.seller)
        created = self.client.post(reverse("order_create"), {
            "client": "Cliente A", "required_date": date.today().isoformat(),
            "product": ["Alitas"], "quantity": ["5"], "unit": ["kg"],
            "reason": ["Pedido manual por rotación"],
        })
        self.assertEqual(created.status_code, 302)
        order = OrderRequest.objects.get(seller=self.seller)
        line = order.lines.get()
        self._upload_sales(sales_book(first, 40, multiplier=3))
        order.refresh_from_db()
        self.assertEqual(order.status, OrderRequest.Status.PENDING)
        self.client.force_login(self.admin)
        review = self.client.post(reverse("order_review", args=[order.pk]), {
            "line_id": line.pk, "decision": "approve", "approved_quantity": "5",
        })
        self.assertEqual(review.status_code, 403)
        revalidated = self.client.post(reverse("order_revalidate", args=[order.pk]), {
            "confirm_revalidation": "1",
        })
        self.assertEqual(revalidated.status_code, 302)
        order.refresh_from_db()
        self.assertEqual(order.revalidated_against, SalesDataset.objects.get(active=True))
        review = self.client.post(reverse("order_review", args=[order.pk]), {
            "line_id": line.pk, "decision": "approve", "approved_quantity": "5",
        })
        self.assertEqual(review.status_code, 302)
        line.refresh_from_db()
        self.assertEqual(line.status, OrderLine.Status.APPROVED)



    def test_admin_can_reset_seller_password_and_toggle_access(self):
        self.client.force_login(self.admin)
        page = self.client.get(reverse("users"))
        self.assertContains(page, "Restablecer contraseñas")
        changed = self.client.post(reverse("users"), {
            "action": "reset_password", "user_id": self.seller.pk,
            "password": "NuevaClaveDelVendedor2026",
        })
        self.assertEqual(changed.status_code, 302)
        self.seller.refresh_from_db()
        self.assertTrue(self.seller.check_password("NuevaClaveDelVendedor2026"))
        self.assertFalse(self.seller.check_password("ClaveDelVendedor123"))
        disabled = self.client.post(reverse("users"), {
            "action": "toggle", "user_id": self.seller.pk,
        })
        self.assertEqual(disabled.status_code, 302)
        self.seller.refresh_from_db()
        self.assertFalse(self.seller.is_active)
        self.assertFalse(self.client.login(
            username=self.seller.username, password="NuevaClaveDelVendedor2026"
        ))

    def test_form_can_submit_one_selected_product_with_other_products_empty(self):
        self._upload_sales(sales_book(date(2025, 1, 1), 10))
        self._assign_seller()
        self.client.force_login(self.seller)
        response = self.client.post(reverse("order_create"), {
            "client": "Cliente A", "required_date": date.today().isoformat(),
            "product": ["Alitas", "Chorizo", ""],
            "quantity": ["4", "", ""],
            "unit": ["kg", "", ""],
            "reason": ["Pedido manual", "", ""],
        })
        self.assertEqual(response.status_code, 302)
        order = OrderRequest.objects.get(seller=self.seller)
        self.assertEqual(list(order.lines.values_list("product", flat=True)), ["Alitas"])
        self.assertContains(self.client.get(reverse("order_detail", args=[order.pk])), "Venta #")
        self.client.force_login(self.admin)
        admin_detail = self.client.get(reverse("order_detail", args=[order.pk]))
        self.assertContains(admin_detail, "Guardar")
        self.assertContains(admin_detail, 'name="approved_quantity" value="4.000"')

    def test_partial_new_product_is_rejected_instead_of_silently_dropped(self):
        self._upload_sales(sales_book(date(2025, 1, 1), 10))
        self._assign_seller()
        self.client.force_login(self.seller)
        response = self.client.post(reverse("order_create"), {
            "client": "Cliente A", "required_date": date.today().isoformat(),
            "product": ["Alitas", ""], "quantity": ["4", "2"],
            "unit": ["kg", "kg"], "reason": ["Conteo manual", ""],
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Complete producto y cantidad")
        self.assertFalse(OrderRequest.objects.exists())


class LegacyMigrationTests(TestCase):
    def test_old_accounts_import_read_only_and_passwords_work(self):
        from hashlib import pbkdf2_hmac
        from pathlib import Path
        from django.contrib.auth.hashers import check_password
        from django.core.management import call_command
        import os
        import sqlite3

        with TemporaryDirectory() as folder:
            path = Path(folder) / "legacy.sqlite3"
            salt = os.urandom(16)
            password = "ClaveAnteriorFirme123"
            digest = pbkdf2_hmac("sha256", password.encode(), salt, 600_000)
            from contextlib import closing
            with closing(sqlite3.connect(path)) as db:
                db.execute("CREATE TABLE users(id INTEGER, username TEXT, salt BLOB, password_hash BLOB, role TEXT, active INTEGER)")
                db.execute("CREATE TABLE assignments(user_id INTEGER, client TEXT)")
                db.execute("INSERT INTO users VALUES(1,?,?,?,?,1)",
                           ("jefe", salt, digest, "admin"))
                db.execute("INSERT INTO assignments VALUES(1,'Cliente A')")
                db.commit()
            before = path.read_bytes()
            call_command("import_legacy_local", path=path, verbosity=0)
            account = get_user_model().objects.get(username="jefe")
            self.assertTrue(account.is_staff)
            self.assertTrue(check_password(password, account.password))
            self.assertTrue(Assignment.objects.filter(user=account, client="Cliente A").exists())
            self.assertEqual(path.read_bytes(), before)
